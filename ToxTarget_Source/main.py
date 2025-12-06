# -*- coding: utf-8 -*-
"""
Main Program Module - Two-Stage Screening UI
"""

import sys
import json
import traceback
import os
import subprocess
import platform
import logging
import shutil
import csv
import asyncio
import aiohttp
import html
import re
from datetime import datetime
from typing import Dict, List, Optional, Tuple
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    logging.warning("openpyxl library not found. Excel export will be disabled.")


from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QLabel, QLineEdit, QTextEdit, QPushButton, QComboBox, QCheckBox,
    QSpinBox, QGroupBox, QScrollArea, QProgressBar, QDialog,
    QDialogButtonBox, QTabWidget, QTableWidget, QTableWidgetItem,
    QHeaderView, QMessageBox, QFileDialog, QFrame, QRadioButton,
    QButtonGroup, QGridLayout, QSplitter, QSizePolicy
)
from PyQt5.QtCore import Qt, pyqtSlot, QObject, pyqtSignal, QThread
from PyQt5.QtGui import QFont, QColor

from core_engine import TwoStageAnalysisEngine, AnalysisThread
from services import AIService
from utils import LiteratureManager, SessionManager, setup_logging, sanitize_filename, DetailedWordExporter
from meeting_manager import ExpertSystemHelper

class ApiConnectionTestWorker(QObject):
    """Worker for API connection test."""
    finished = pyqtSignal(bool, str)

    def __init__(self, proxy_url: Optional[str], model_name: str, api_key: str):
        super().__init__()
        self.proxy_url = proxy_url
        self.model_name = model_name
        self.api_key = api_key
        self.ai_service_tester = AIService(model_name=model_name, api_key=api_key)

    async def _test_async(self) -> Tuple[bool, str]:
        """Async API connection test coroutine."""
        try:
            success, message = await AIService.test_connection(
                 model_name=self.model_name,
                 api_key=self.api_key,
                 proxy_url=self.proxy_url
            )
            return success, message
        except Exception as e:
            logging.error(f"API connection test worker failed unexpectedly: {e}", exc_info=True)
            return False, f"An unexpected error occurred during the test: {e}"

    @pyqtSlot()
    def run(self):
        """Runs the async test."""
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        try:
            success, message = loop.run_until_complete(self._test_async())
            self.finished.emit(success, message)
        finally:
            loop.close()


class ExperimentalDataDialog(QDialog):
    """Experimental Data Input Dialog"""

    def __init__(self, current_data: str = '', parent=None):
        super().__init__(parent)
        self.setWindowTitle('Add Experimental Data')
        self.setModal(True)
        self.resize(600, 400)

        layout = QVBoxLayout(self)

        info_label = QLabel(
            'Enter your experimental observations, preliminary data, or supplementary information. '
            'This will enhance AI analysis accuracy.'
        )
        info_label.setWordWrap(True)
        layout.addWidget(info_label)

        self.text_edit = QTextEdit()
        self.text_edit.setPlaceholderText(
            'Example:\n'
            '- Western Blot shows ALB expression decreased ~40% after PFOA treatment\n'
            '- qPCR: CYP1A1 mRNA increased 2.5-fold\n'
            '- Cell viability: IC50 approximately 100µM'
        )
        self.text_edit.setPlainText(current_data)
        self.text_edit.setStyleSheet('background-color: white; color: black;')
        layout.addWidget(self.text_edit)

        button_box = QDialogButtonBox(
            QDialogButtonBox.Save | QDialogButtonBox.Cancel
        )
        button_box.accepted.connect(self.accept)
        button_box.rejected.connect(self.reject)
        layout.addWidget(button_box)

    def get_data(self) -> str:
        return self.text_edit.toPlainText().strip()


class SaveableTextDialog(QDialog):
    """
    A reusable dialog to display text content (Markdown or plain text),
    with a built-in "Save As..." button.
    """
    def __init__(self, title: str, content: str, default_filename: str = "export.txt", is_markdown: bool = True, parent=None):
        super().__init__(parent)
        self.setWindowTitle(title)
        self.setModal(True)
        self.resize(1000, 700)
        
        self.content_to_save = content
        self.default_filename = default_filename
        
        layout = QVBoxLayout(self)
        
        self.text_browser = QTextEdit()
        self.text_browser.setReadOnly(True)
        self.text_browser.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        
        if is_markdown:
            # Attempt to render Markdown
            self.text_browser.setMarkdown(self.content_to_save)
        else:
            # Show plain text
            self.text_browser.setPlainText(self.content_to_save)
            
        layout.addWidget(self.text_browser)
        
        # Buttons
        button_box = QDialogButtonBox()
        save_btn = button_box.addButton("Save As...", QDialogButtonBox.AcceptRole)
        close_btn = button_box.addButton("Close", QDialogButtonBox.RejectRole)
        
        save_btn.clicked.connect(self._save_content)
        close_btn.clicked.connect(self.reject)
        
        layout.addWidget(button_box)
        
    def _save_content(self):
        """Open file save dialog and save content"""
        
        # Determine file filter based on content type
        if self.default_filename.endswith('.txt'):
            filter_str = "Text Files (*.txt);;All Files (*)"
        elif self.default_filename.endswith('.md'):
            filter_str = "Markdown Files (*.md);;Text Files (*.txt);;All Files (*)"
        else:
            filter_str = "All Files (*)"

        filename, _ = QFileDialog.getSaveFileName(
            self, 'Save As', self.default_filename, filter_str
        )
        
        if filename:
            try:
                content_to_write = self.content_to_save
                
                # Simple conversion if saving Markdown as .txt
                if filename.endswith('.txt') and self.content_to_save.startswith("#"):
                     content_to_write = re.sub(r'###\s*(.*)', r'\n--- \1 ---', self.content_to_save)
                     content_to_write = re.sub(r'\*\*(.*?)\*\*', r'\1', content_to_write)
                     content_to_write = re.sub(r'^\s*-\s*', r'  * ', content_to_write, flags=re.MULTILINE)
                
                with open(filename, 'w', encoding='utf-8') as f:
                    f.write(content_to_write)
                QMessageBox.information(self, 'Success', f'File saved to:\n{filename}')
            except Exception as e:
                QMessageBox.critical(self, 'Error', f'File save failed:\n{str(e)}')


class QuickScreeningResultDialog(QDialog):
    """Quick Screening Results Display Dialog"""

    def __init__(self, screening_result: Dict, session_inputs: Dict, parent=None):
        super().__init__(parent)
        self.screening_result = screening_result
        self.session_inputs = session_inputs
        self.selected_proteins = []

        self.setWindowTitle('Quick Screening Results - Top Recommended Proteins')
        self.setModal(True)
        self.resize(1000, 700)

        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout(self)

        title = QLabel('Quick Screening Results')
        title.setStyleSheet('font-size: 18pt; font-weight: bold; color: #2c3e50;')
        title.setAlignment(Qt.AlignCenter)
        layout.addWidget(title)

        stats = self.screening_result
        stats_widget = QFrame()
        stats_widget.setStyleSheet('background-color: #e8f4f8; padding: 10px; border-radius: 5px;')
        stats_layout = QHBoxLayout(stats_widget)

        stats_layout.addWidget(QLabel(f"Total Evaluated: {stats.get('total_evaluated', 'N/A')}"))

        score_dist = stats.get('score_distribution', {})
        if score_dist:
            high_threshold = 70
            medium_threshold = 50
            stats_layout.addWidget(QLabel(
                f"High Score(>={high_threshold}): {score_dist.get('high', 0)} | "
                f"Medium({medium_threshold}-{high_threshold}): {score_dist.get('medium', 0)} | "
                f"Low(<{medium_threshold}): {score_dist.get('low', 0)}"
            ))
        else:
            stats_layout.addWidget(QLabel("Score distribution unavailable"))

        stats_layout.addStretch()
        layout.addWidget(stats_widget)

        info_label = QLabel(
            'Top candidate proteins after AI review.\n'
            'Select proteins to proceed to detailed screening.'
        )
        info_label.setWordWrap(True)
        info_label.setStyleSheet('padding: 10px; color: #555; font-size: 10pt;')
        layout.addWidget(info_label)

        table = QTableWidget()
        table.setColumnCount(7) # Increased column count from 6 to 7
        table.setHorizontalHeaderLabels(['Select', 'Rank', 'Protein', 'Final Score', 'Base Score', 'Affinity Bonus', 'Scores per Round'])

        top_list = self.screening_result.get('top10', [])
        table.setRowCount(len(top_list))

        self.checkboxes = []

        for row, item in enumerate(top_list):
            checkbox = QCheckBox()
            checkbox.setChecked(True)
            self.checkboxes.append(checkbox)
            table.setCellWidget(row, 0, checkbox)

            rank_item = QTableWidgetItem(str(item.get('rank', row + 1)))
            rank_item.setTextAlignment(Qt.AlignCenter)
            table.setItem(row, 1, rank_item)

            protein_item = QTableWidgetItem(item.get('protein', ''))
            table.setItem(row, 2, protein_item)

            score_val = item.get('final_score', 0)
            score_item = QTableWidgetItem(f"{score_val:.2f}")
            score_item.setTextAlignment(Qt.AlignCenter)
            if score_val >= 80:
                score_item.setBackground(QColor(212, 237, 218))
            elif score_val >= 60:
                score_item.setBackground(QColor(255, 243, 205))
            table.setItem(row, 3, score_item)

            base_score_val = item.get('base_score')
            base_score_str = f"{base_score_val:.2f}" if base_score_val is not None else "N/A"
            base_score_item = QTableWidgetItem(base_score_str)
            base_score_item.setTextAlignment(Qt.AlignCenter)
            table.setItem(row, 4, base_score_item)

            bonus_val = item.get('binding_bonus')
            bonus_str = f"+{bonus_val}" if bonus_val is not None else "N/A"
            bonus_item = QTableWidgetItem(bonus_str)
            bonus_item.setTextAlignment(Qt.AlignCenter)
            table.setItem(row, 5, bonus_item)

            rounds_val = item.get('scores_per_round')
            rounds_str = ', '.join(map(str, rounds_val)) if rounds_val and isinstance(rounds_val, list) else "-"
            rounds_item = QTableWidgetItem(rounds_str)
            rounds_item.setToolTip(f"Scores from {len(rounds_val) if isinstance(rounds_val, list) else 0} rounds of evaluation")
            rounds_item.setTextAlignment(Qt.AlignCenter)
            table.setItem(row, 6, rounds_item)


        table.setSelectionBehavior(QTableWidget.SelectRows)
        table.horizontalHeader().setStretchLastSection(True)
        table.setAlternatingRowColors(True)
        table.verticalHeader().setVisible(False)
        table.setEditTriggers(QTableWidget.NoEditTriggers)

        table.setColumnWidth(0, 50)  # Select
        table.setColumnWidth(1, 50)  # Rank
        table.setColumnWidth(2, 150) # Protein
        table.setColumnWidth(3, 80)  # Final Score
        table.setColumnWidth(4, 80)  # Base Score
        table.setColumnWidth(5, 90)  # Affinity Bonus
        table.setColumnWidth(6, 120) # Scores per Round (new)

        layout.addWidget(table)

        detail_text = QTextEdit()
        detail_text.setReadOnly(True)
        detail_text.setMaximumHeight(150)
        detail_text.setStyleSheet('background-color: #f8f9fa; border: 1px solid #ddd; color: black;')
        layout.addWidget(QLabel('Detailed Rationale (for selected row):'))
        layout.addWidget(detail_text)

        table.cellClicked.connect(lambda row, col: self._show_details(row, detail_text))

        note_label = QLabel(
            'Note: Only proteins selected above will proceed to detailed screening (literature search + deep review).'
        )
        note_label.setStyleSheet('color: #d32f2f; font-weight: bold; padding: 10px;')
        note_label.setWordWrap(True)
        layout.addWidget(note_label)

        button_layout = QHBoxLayout()

        select_all_btn = QPushButton('Select All')
        select_all_btn.clicked.connect(lambda: self._select_all(True))
        button_layout.addWidget(select_all_btn)

        deselect_all_btn = QPushButton('Deselect All')
        deselect_all_btn.clicked.connect(lambda: self._select_all(False))
        button_layout.addWidget(deselect_all_btn)

        button_layout.addStretch()

        export_btn = QPushButton('Export Results')
        export_btn.clicked.connect(self._export_results)
        button_layout.addWidget(export_btn)

        view_full_btn = QPushButton('View Full Ranking')
        view_full_btn.clicked.connect(self._view_full_ranking)
        button_layout.addWidget(view_full_btn)

        proceed_btn = QPushButton('Proceed to Detailed Screening')
        proceed_btn.setStyleSheet(
            'background-color: #4CAF50; color: white; padding: 10px 30px; '
            'font-weight: bold; font-size: 12pt;'
        )
        proceed_btn.clicked.connect(self.accept)
        button_layout.addWidget(proceed_btn)

        skip_btn = QPushButton('Skip (Complete Quick Mode)')
        skip_btn.clicked.connect(self.reject)
        button_layout.addWidget(skip_btn)

        layout.addLayout(button_layout)

    def _show_details(self, row: int, detail_widget: QTextEdit):
        top_list = self.screening_result.get('top10', [])
        if row < len(top_list):
            item = top_list[row]
            details = f"<h3>{item.get('protein', 'N/A')} - Rank #{item.get('rank', '')}</h3>"
            details += f"<p><b>Final Score:</b> {item.get('final_score', 0):.2f}</p>"
            details += f"<p><b>Rationale & Score Breakdown:</b> {item.get('rationale', 'N/A')}</p>"
            detail_widget.setHtml(details)

    def _select_all(self, checked: bool):
        for cb in self.checkboxes:
            cb.setChecked(checked)

    def get_selected_proteins(self) -> List[str]:
        top_list = self.screening_result.get('top10', [])
        selected = []
        for i, cb in enumerate(self.checkboxes):
            if cb.isChecked() and i < len(top_list):
                selected.append(top_list[i].get('protein', ''))
        return selected

    def _export_results(self):
        file_path, selected_filter = QFileDialog.getSaveFileName(
            self,
            'Export Quick Screening Results',
            f'quick_results_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx',
            'Excel Files (*.xlsx);;CSV Files (*.csv)'
        )

        if not file_path:
            return

        try:
            top_list = self.screening_result.get('top10', [])
            if 'xlsx' in selected_filter:
                if not OPENPYXL_AVAILABLE:
                    raise ImportError("openpyxl library is required for Excel export.")
                self._export_excel(top_list, file_path)
            else:
                self._export_csv(top_list, file_path)
            QMessageBox.information(self, 'Success', f'Results exported to:\n{file_path}')
        except ImportError as ie:
            QMessageBox.critical(self, 'Export Failed', f'Missing library: {str(ie)}')
        except Exception as e:
            QMessageBox.critical(self, 'Export Failed', f'Error: {str(e)}')

    def _export_excel(self, items: List[Dict], file_path: str):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Quick Screening Results'

        headers = ['Rank', 'Protein', 'Final Score', 'Base Score', 'Affinity Bonus', 'Scores per Round', 'Rationale']
        ws.append(headers)

        header_font = Font(bold=True, size=12, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_align = Alignment(horizontal='center', vertical='center')
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(1, col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        for item in items:
            rounds_val = item.get('scores_per_round')
            rounds_str = ', '.join(map(str, rounds_val)) if rounds_val and isinstance(rounds_val, list) else ""
            ws.append([
                item.get('rank', ''),
                item.get('protein', ''),
                f"{item.get('final_score', 0):.2f}",
                f"{item.get('base_score', 0):.2f}" if item.get('base_score') is not None else "N/A",
                f"{item.get('binding_bonus', 0)}" if item.get('binding_bonus') is not None else "N/A",
                rounds_str, # Include the scores per round string
                item.get('rationale', '')
            ])

        # Adjust column widths (7 columns)
        for col_idx, col_width in enumerate([8, 20, 12, 12, 15, 18, 60], start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = col_width

        wb.save(file_path)

    def _export_csv(self, items: List[Dict], file_path: str):
        with open(file_path, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            writer.writerow(['Rank', 'Protein', 'Final Score', 'Base Score', 'Affinity Bonus', 'Scores per Round', 'Rationale'])
            for item in items:
                rounds_val = item.get('scores_per_round')
                rounds_str = ', '.join(map(str, rounds_val)) if rounds_val and isinstance(rounds_val, list) else ""
                writer.writerow([
                    item.get('rank', ''),
                    item.get('protein', ''),
                    f"{item.get('final_score', 0):.2f}",
                    f"{item.get('base_score', 0):.2f}" if item.get('base_score') is not None else "N/A",
                    f"{item.get('binding_bonus', 0)}" if item.get('binding_bonus') is not None else "N/A",
                    rounds_str, # Include the scores per round string
                    item.get('rationale', '')
                ])


    def _view_full_ranking(self):
        full_ranking = self.screening_result.get('full_ranking', [])
        if not full_ranking:
            QMessageBox.information(self, 'No Data', 'Full ranking data not available.')
            return

        dialog = FullRankingDialog(full_ranking, self.session_inputs, self)
        dialog.exec_()


class FullRankingDialog(QDialog):
    """Dialog to display and search full ranking."""

    def __init__(self, full_ranking: List[Dict], session_inputs: Dict, parent=None):
        super().__init__(parent)
        self.full_ranking = full_ranking
        self.session_inputs = session_inputs

        self.setWindowTitle(f'Full Protein Ranking - All {len(full_ranking)} Evaluated Proteins')
        self.setModal(True)
        self.resize(1200, 800)

        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout(self)

        title = QLabel('Complete Protein Ranking')
        title.setStyleSheet('font-size: 18pt; font-weight: bold; color: #2c3e50;')
        title.setAlignment(Qt.AlignCenter)
        layout.addWidget(title)

        filter_widget = QFrame()
        filter_widget.setStyleSheet('background-color: #f0f0f0; padding: 10px; border-radius: 5px;')
        filter_layout = QHBoxLayout(filter_widget)

        filter_layout.addWidget(QLabel('Search:'))
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText('Enter protein name...')
        self.search_input.textChanged.connect(self._filter_table)
        filter_layout.addWidget(self.search_input)

        filter_layout.addWidget(QLabel('Min Score:'))
        self.score_filter = QComboBox()
        self.score_filter.addItems(['All', '>=90', '>=80', '>=70', '>=60', '>=50', '>=40', '>=30']) # Adjusted for 100+ scale
        self.score_filter.currentTextChanged.connect(self._filter_table)
        filter_layout.addWidget(self.score_filter)

        filter_layout.addWidget(QLabel('Show:'))
        self.topn_filter = QComboBox()
        self.topn_filter.addItems(['All', 'Top 10', 'Top 20', 'Top 50', 'Top 100'])
        self.topn_filter.currentTextChanged.connect(self._filter_table)
        filter_layout.addWidget(self.topn_filter)

        filter_layout.addStretch()
        layout.addWidget(filter_widget)

        self.table = QTableWidget()
        self.table.setColumnCount(4) # Keep 4 columns for full view for simplicity
        self.table.setHorizontalHeaderLabels(['Rank', 'Protein', 'Final Score', 'In Top 10'])
        self.table.horizontalHeader().setStretchLastSection(True)
        self.table.setAlternatingRowColors(True)
        self.table.verticalHeader().setVisible(False)
        self.table.setSelectionBehavior(QTableWidget.SelectRows)
        self.table.setEditTriggers(QTableWidget.NoEditTriggers)

        self.table.setColumnWidth(0, 80)
        self.table.setColumnWidth(1, 200)
        self.table.setColumnWidth(2, 100)
        self.table.setColumnWidth(3, 100)

        layout.addWidget(self.table)

        self.detail_text = QTextEdit()
        self.detail_text.setReadOnly(True)
        self.detail_text.setMaximumHeight(150)
        self.detail_text.setStyleSheet('background-color: #f8f9fa; border: 1px solid #ddd; color: black;')
        layout.addWidget(QLabel('Detailed Rationale (for selected protein):'))
        layout.addWidget(self.detail_text)

        self.table.cellClicked.connect(lambda row, col: self._show_details(row))

        button_layout = QHBoxLayout()
        export_btn = QPushButton('Export This View')
        export_btn.clicked.connect(self._export_view)
        button_layout.addWidget(export_btn)
        button_layout.addStretch()
        close_btn = QPushButton('Close')
        close_btn.clicked.connect(self.accept)
        button_layout.addWidget(close_btn)
        layout.addLayout(button_layout)

        self._populate_table()

    def _populate_table(self, items=None):
        if items is None:
            items = self.full_ranking

        self.table.setRowCount(len(items))

        for row, item in enumerate(items):
            rank_item = QTableWidgetItem(str(item.get('rank', '')))
            rank_item.setTextAlignment(Qt.AlignCenter)
            self.table.setItem(row, 0, rank_item)

            protein_item = QTableWidgetItem(item.get('protein', ''))
            self.table.setItem(row, 1, protein_item)

            score_val = item.get('final_score', 0)
            score_item = QTableWidgetItem(f"{score_val:.2f}")
            score_item.setTextAlignment(Qt.AlignCenter)
            if score_val >= 80:
                score_item.setBackground(QColor(212, 237, 218))
            elif score_val >= 60:
                score_item.setBackground(QColor(255, 243, 205))
            self.table.setItem(row, 2, score_item)

            top10_item = QTableWidgetItem('Yes' if item.get('in_top10', False) else 'No')
            top10_item.setTextAlignment(Qt.AlignCenter)
            self.table.setItem(row, 3, top10_item)

    def _filter_table(self):
        search_text = self.search_input.text().lower()
        score_filter = self.score_filter.currentText()
        topn_filter = self.topn_filter.currentText()

        filtered = self.full_ranking

        if search_text:
            filtered = [item for item in filtered if search_text in item.get('protein', '').lower()]

        if score_filter != 'All':
            min_score = float(score_filter.replace('>=', ''))
            filtered = [item for item in filtered if item.get('final_score', 0) >= min_score]

        if topn_filter != 'All':
            n = int(topn_filter.split()[1])
            # Ensure sorting before slicing for top N
            filtered.sort(key=lambda x: x.get('final_score', 0), reverse=True)
            filtered = filtered[:n]

        self._populate_table(filtered)

    def _show_details(self, row: int):
        current_items = []
        for r in range(self.table.rowCount()):
            protein = self.table.item(r, 1).text() if self.table.item(r, 1) else ''
            for item in self.full_ranking:
                if item.get('protein', '') == protein:
                    current_items.append(item)
                    break

        if row < len(current_items):
            item = current_items[row]
            details = f"<h3>{item.get('protein', 'N/A')} - Rank #{item.get('rank', '')}</h3>"
            details += f"<p><b>Final Score:</b> {item.get('final_score', 0):.2f}</p>"
            details += f"<p><b>Rationale & Score Breakdown:</b> {item.get('rationale', 'N/A')}</p>"
            self.detail_text.setHtml(details)

    def _export_view(self):
        current_items = []
        for r in range(self.table.rowCount()):
            protein = self.table.item(r, 1).text() if self.table.item(r, 1) else ''
            for item in self.full_ranking:
                if item.get('protein', '') == protein:
                    current_items.append(item)
                    break

        if not current_items:
            QMessageBox.warning(self, 'No Data', 'No data to export.')
            return

        file_path, selected_filter = QFileDialog.getSaveFileName(
            self,
            'Export Current View',
            f'filtered_ranking_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx',
            'Excel Files (*.xlsx);;CSV Files (*.csv)'
        )

        if not file_path:
            return

        try:
            if 'xlsx' in selected_filter:
                if not OPENPYXL_AVAILABLE:
                    raise ImportError("openpyxl library is required for Excel export.")
                self._export_excel(current_items, file_path)
            else:
                self._export_csv(current_items, file_path)
            QMessageBox.information(self, 'Success', f'View exported to:\n{file_path}')
        except ImportError as ie:
            QMessageBox.critical(self, 'Export Failed', f'Missing library: {str(ie)}')
        except Exception as e:
            QMessageBox.critical(self, 'Export Failed', f'Error: {str(e)}')

    def _export_excel(self, items: List[Dict], file_path: str):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Filtered Ranking'

        headers = ['Rank', 'Protein', 'Final Score', 'Base Score', 'Affinity Bonus', 'Scores per Round', 'In Top 10', 'Rationale']
        ws.append(headers)

        header_font = Font(bold=True, size=12, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_align = Alignment(horizontal='center', vertical='center')
        for col_idx, header_text in enumerate(headers, 1):
            cell = ws.cell(1, col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        for item in items:
            rounds_val = item.get('scores_per_round')
            rounds_str = ', '.join(map(str, rounds_val)) if rounds_val and isinstance(rounds_val, list) else ""
            ws.append([
                item.get('rank', ''),
                item.get('protein', ''),
                f"{item.get('final_score', 0):.2f}",
                f"{item.get('base_score', 0):.2f}" if item.get('base_score') is not None else "N/A",
                f"{item.get('binding_bonus', 0)}" if item.get('binding_bonus') is not None else "N/A",
                rounds_str, # Include scores per round
                'Yes' if item.get('in_top10', False) else 'No',
                item.get('rationale', '')
            ])

        # Adjust column widths (8 columns)
        for col_idx, col_width in enumerate([8, 20, 12, 12, 15, 18, 15, 60], start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = col_width

        wb.save(file_path)

    def _export_csv(self, items: List[Dict], file_path: str):
        with open(file_path, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            writer.writerow(['Rank', 'Protein', 'Final Score', 'Base Score', 'Affinity Bonus', 'Scores per Round', 'In Top 10', 'Rationale'])
            for item in items:
                rounds_val = item.get('scores_per_round')
                rounds_str = ', '.join(map(str, rounds_val)) if rounds_val and isinstance(rounds_val, list) else ""
                writer.writerow([
                    item.get('rank', ''),
                    item.get('protein', ''),
                    f"{item.get('final_score', 0):.2f}",
                    f"{item.get('base_score', 0):.2f}" if item.get('base_score') is not None else "N/A",
                    f"{item.get('binding_bonus', 0)}" if item.get('binding_bonus') is not None else "N/A",
                    rounds_str, # Include scores per round
                    'Yes' if item.get('in_top10', False) else 'No',
                    item.get('rationale', '')
                ])


class LiteratureReviewDialog(QDialog):
    """Literature Review Selection Dialog"""

    def __init__(self, literature_data: Dict, session_inputs: Dict, parent=None):
        super().__init__(parent)
        self.literature_data = literature_data
        self.session_inputs = session_inputs
        self.checkboxes = {}

        self.setWindowTitle('Literature Review - Select Articles for Analysis')
        self.setModal(True)
        self.resize(1400, 800)

        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout(self)

        title = QLabel('Literature Search Results')
        title.setStyleSheet('font-size: 18pt; font-weight: bold; color: #2c3e50;')
        title.setAlignment(Qt.AlignCenter)
        layout.addWidget(title)

        info_label = QLabel(
            'AI has retrieved relevant articles for each protein. '
            'Review and select articles to include in detailed analysis. '
            'All articles are pre-selected by default.'
        )
        info_label.setWordWrap(True)
        info_label.setStyleSheet('padding: 10px; color: #555;')
        layout.addWidget(info_label)

        tabs = QTabWidget()

        for protein, data in sorted(self.literature_data.items()):
            tab = self._create_protein_tab(protein, data)
            tabs.addTab(tab, protein)

        layout.addWidget(tabs)

        button_layout = QHBoxLayout()

        select_all_btn = QPushButton('Select All Literature')
        select_all_btn.clicked.connect(lambda: self._select_all_literature(True))
        button_layout.addWidget(select_all_btn)

        deselect_all_btn = QPushButton('Deselect All')
        deselect_all_btn.clicked.connect(lambda: self._select_all_literature(False))
        button_layout.addWidget(deselect_all_btn)

        button_layout.addStretch()

        proceed_btn = QPushButton('Proceed to Analysis')
        proceed_btn.setStyleSheet(
            'background-color: #5c6bc0; color: white; padding: 10px 30px; '
            'font-weight: bold; font-size: 12pt;'
        )
        proceed_btn.clicked.connect(self.accept)
        button_layout.addWidget(proceed_btn)

        cancel_btn = QPushButton('Cancel')
        cancel_btn.clicked.connect(self.reject)
        button_layout.addWidget(cancel_btn)

        layout.addLayout(button_layout)

    def _create_protein_tab(self, protein: str, data: Dict) -> QWidget:
        tab = QWidget()
        layout = QVBoxLayout(tab)

        stats_text = f"<b>Literature Retrieved:</b> {len(data.get('articles', []))} articles"
        strategies = data.get('search_strategies')
        if isinstance(strategies, (list, tuple, dict)):
             stats_text += f" | <b>Search Strategies Used:</b> {len(strategies)}"
        elif strategies is not None:
             logging.warning(f"Unexpected type for 'search_strategies' for {protein}: {type(strategies)}")

        quality_metrics = data.get('quality_metrics', {})
        if isinstance(quality_metrics, dict):
            avg_rel = quality_metrics.get('avg_relevance_score', 0)
            high_qual = quality_metrics.get('high_quality_count', 0)
            stats_text += f" | <b>Avg. Relevance:</b> {avg_rel:.1f}/10 | <b>High Quality (>=7):</b> {high_qual}"
        else:
            logging.warning(f"Unexpected type for 'quality_metrics' for {protein}: {type(quality_metrics)}")

        stats_label = QLabel(stats_text)
        stats_label.setStyleSheet('padding: 8px; background-color: #e8f5e9; border-radius: 4px;')
        layout.addWidget(stats_label)

        table = QTableWidget()
        table.setColumnCount(6)
        table.setHorizontalHeaderLabels(['Select', 'PMID', 'Title', 'Year', 'Relevance', 'Strategy'])

        articles = data.get('articles', [])
        table.setRowCount(len(articles))

        self.checkboxes[protein] = []

        for row, article in enumerate(articles):
            checkbox = QCheckBox()
            checkbox.setChecked(True)
            self.checkboxes[protein].append(checkbox)
            table.setCellWidget(row, 0, checkbox)

            pmid_item = QTableWidgetItem(article.get('pmid', 'N/A'))
            table.setItem(row, 1, pmid_item)

            title_text = article.get('title', 'N/A')
            title_item = QTableWidgetItem(title_text[:60] + ('...' if len(title_text) > 60 else ''))
            title_item.setToolTip(title_text)
            table.setItem(row, 2, title_item)

            year_item = QTableWidgetItem(str(article.get('year', 'N/A')))
            year_item.setTextAlignment(Qt.AlignCenter)
            table.setItem(row, 3, year_item)

            score = article.get('ai_relevance_score', 0)
            score_item = QTableWidgetItem(f"{score:.1f}/10")
            score_item.setTextAlignment(Qt.AlignCenter)
            if score >= 7: score_item.setBackground(QColor(212, 237, 218))
            elif score >= 5: score_item.setBackground(QColor(255, 243, 205))
            else: score_item.setBackground(QColor(255, 235, 238))
            table.setItem(row, 4, score_item)

            strategy_item = QTableWidgetItem(article.get('search_strategy', 'N/A'))
            table.setItem(row, 5, strategy_item)

        table.setSelectionBehavior(QTableWidget.SelectRows)
        table.horizontalHeader().setStretchLastSection(True)
        table.setAlternatingRowColors(True)
        table.verticalHeader().setVisible(False)
        table.setEditTriggers(QTableWidget.NoEditTriggers)

        table.setColumnWidth(0, 60)
        table.setColumnWidth(1, 100)
        table.setColumnWidth(2, 400)
        table.setColumnWidth(3, 80)
        table.setColumnWidth(4, 80)
        table.setColumnWidth(5, 150)

        layout.addWidget(table)

        return tab

    def _select_all_literature(self, checked: bool):
        for protein_checkboxes in self.checkboxes.values():
            for cb in protein_checkboxes:
                cb.setChecked(checked)

    def get_selected_literature(self) -> Dict:
        selected = {}

        for protein, data in self.literature_data.items():
            articles = data.get('articles', [])
            checkboxes_list = self.checkboxes.get(protein, [])

            selected_articles = []
            for i in range(len(articles)):
                if i < len(checkboxes_list) and checkboxes_list[i].isChecked():
                    selected_articles.append(articles[i])

            if selected_articles:
                selected_data = {k: v for k, v in data.items() if k != 'articles'} if data else {}
                selected_data['articles'] = selected_articles
                selected[protein] = selected_data

        return selected


class TwoStageMainWindow(QMainWindow):
    """Main Window"""

    def __init__(self):
        super().__init__()

        self.setWindowTitle('ToxTarget-AI: AI-Driven Target Protein Screening Platform')
        self.setGeometry(100, 100, 1400, 900)

        self.engine = TwoStageAnalysisEngine()
        self.analysis_thread = None
        self.experimental_data = ''
        self.session_manager = SessionManager()
        self.literature_manager = LiteratureManager()

        self.engine.log_signal.connect(self.add_log)
        self.engine.detailed_log_signal.connect(self.add_detailed_log)
        self.engine.progress_signal.connect(self.update_progress)
        self.engine.quick_screening_complete_signal.connect(self.on_quick_screening_complete)
        self.engine.literature_ready_signal.connect(self.show_literature_review)
        self.engine.analysis_complete_signal.connect(self.on_analysis_complete)
        self.engine.error_signal.connect(self.on_error)

        self.log_buffer = []
        self.current_session_id = None
        self.session_data = {}

        self.init_ui()

        self.on_mode_changed()
        initial_mode = "Quick" if self.quick_mode_radio.isChecked() else "Detailed"
        logging.info(f"Application initialized in {initial_mode} Mode")


    def init_ui(self):
        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        main_layout = QVBoxLayout(central_widget)

        splitter = QSplitter(Qt.Horizontal)
        left_panel = self.create_left_panel()
        right_panel = self.create_right_panel()
        splitter.addWidget(left_panel)
        splitter.addWidget(right_panel)
        splitter.setSizes([450, 950])
        main_layout.addWidget(splitter)

    def create_left_panel(self) -> QWidget:
        panel = QWidget()
        layout = QVBoxLayout(panel)

        title = QLabel('Research Parameters Configuration')
        title.setStyleSheet('font-size: 16pt; font-weight: bold; padding: 10px;')
        layout.addWidget(title)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll_widget = QWidget()
        scroll_layout = QVBoxLayout(scroll_widget)

        # Screening Mode Group
        self.mode_group = QGroupBox('Screening Mode')
        mode_layout = QVBoxLayout()
        self.mode_btn_group = QButtonGroup(self)

        self.quick_mode_radio = QRadioButton('Quick Screening (for large candidate sets)')
        self.quick_mode_radio.setChecked(True)
        self.quick_mode_radio.toggled.connect(self.on_mode_changed)
        self.mode_btn_group.addButton(self.quick_mode_radio)
        mode_layout.addWidget(self.quick_mode_radio)
        quick_hint = QLabel('  * AI knowledge base only\n  * Rapid screening for top candidates')
        quick_hint.setStyleSheet('color: gray; font-size: 9pt; margin-left: 20px;')
        mode_layout.addWidget(quick_hint)

        # Quick Mode Options
        self.quick_options_widget = QWidget()
        grid_layout = QGridLayout(self.quick_options_widget)
        grid_layout.setContentsMargins(20, 0, 0, 0)

        grid_layout.addWidget(QLabel('Select Top:'), 0, 0)
        self.top_n_spinbox = QSpinBox()
        self.top_n_spinbox.setRange(5, 50)
        self.top_n_spinbox.setValue(10)
        self.top_n_spinbox.setSuffix(' proteins')
        grid_layout.addWidget(self.top_n_spinbox, 0, 1)

        consistency_label = QLabel('Consistency Rounds:')
        consistency_label.setToolTip('Run evaluation multiple times for more stable results.\n1 = Default, 2-3 = Higher consistency (slower, more API calls).')
        grid_layout.addWidget(consistency_label, 1, 0)
        self.consistency_spinbox = QSpinBox()
        self.consistency_spinbox.setRange(1, 3)
        self.consistency_spinbox.setValue(1)
        self.consistency_spinbox.setSpecialValueText('1 (Default)')
        grid_layout.addWidget(self.consistency_spinbox, 1, 1)

        grid_layout.setColumnStretch(2, 1)
        mode_layout.addWidget(self.quick_options_widget)


        self.detailed_mode_radio = QRadioButton('Detailed Screening (recommended for 10-15 proteins)')
        self.detailed_mode_radio.toggled.connect(self.on_mode_changed)
        self.mode_btn_group.addButton(self.detailed_mode_radio)
        mode_layout.addWidget(self.detailed_mode_radio)
        detailed_hint = QLabel('  * AI + literature search\n  * Comprehensive review with expert simulation')
        detailed_hint.setStyleSheet('color: gray; font-size: 9pt; margin-left: 20px;')
        mode_layout.addWidget(detailed_hint)

        self.mode_group.setLayout(mode_layout)
        scroll_layout.addWidget(self.mode_group)

        # Core Inputs
        scroll_layout.addWidget(QLabel('Pollutant/Compound *'))
        self.pollutant_input = QLineEdit()
        self.pollutant_input.setPlaceholderText('e.g., PFOA, BPA, Lead')
        self.pollutant_input.setStyleSheet('background-color: white; color: black;')
        scroll_layout.addWidget(self.pollutant_input)

        scroll_layout.addWidget(QLabel('Candidate Proteins * (one per line/comma/semicolon)'))
        self.proteins_input = QTextEdit()
        self.proteins_input.setPlaceholderText('Quick: dozens-hundreds\nDetailed: 10-15 recommended\n\nALB\nFABP1, Serum albumin\nHSA; ALBU')
        self.proteins_input.setMaximumHeight(120)
        self.proteins_input.textChanged.connect(self.update_protein_count)
        self.proteins_input.setStyleSheet('background-color: white; color: black;')
        scroll_layout.addWidget(self.proteins_input)
        self.protein_count_label = QLabel('0 proteins')
        self.protein_count_label.setStyleSheet('color: gray; font-size: 9pt;')
        scroll_layout.addWidget(self.protein_count_label)
        self.mode_hint_label = QLabel('')
        self.mode_hint_label.setStyleSheet('color: #ff9800; font-size: 9pt; font-weight: bold;')
        scroll_layout.addWidget(self.mode_hint_label)

        scroll_layout.addWidget(QLabel('Target Toxic Phenotypes (optional, one per line)'))
        self.phenotypes_input = QTextEdit()
        self.phenotypes_input.setPlaceholderText('e.g., hepatotoxicity\nNAFLD\nsteatosis')
        self.phenotypes_input.setMaximumHeight(60)
        self.phenotypes_input.setStyleSheet('background-color: white; color: black;')
        scroll_layout.addWidget(self.phenotypes_input)

        scroll_layout.addWidget(QLabel('Target Tissue *'))
        self.tissue_combo = QComboBox()
        self.tissue_combo.addItems(['Select Tissue', 'Liver', 'Kidney', 'Brain', 'Heart', 'Lung', 'Blood', 'Adipose Tissue', 'Muscle', 'Skin', 'Intestine'])
        scroll_layout.addWidget(self.tissue_combo)

        scroll_layout.addWidget(QLabel('Research Species *'))
        self.species_combo = QComboBox()
        self.species_combo.addItems(['Human (Homo sapiens)', 'Mouse (Mus musculus)', 'Zebrafish (Danio rerio)'])
        self.species_combo.currentIndexChanged.connect(self.on_species_changed)
        scroll_layout.addWidget(self.species_combo)
        self.species_data_hint = QLabel('')
        self.species_data_hint.setStyleSheet('color: #666; font-size: 9pt; font-style: italic;')
        self.species_data_hint.setWordWrap(True)
        scroll_layout.addWidget(self.species_data_hint)
        self.on_species_changed()

        # Email Input
        self.email_label = QLabel('Email Address * (Required for Detailed Mode & PubMed/EPMC)')
        scroll_layout.addWidget(self.email_label)
        self.email_input = QLineEdit()
        self.email_input.setPlaceholderText('your.email@example.com')
        self.email_input.setStyleSheet('background-color: white; color: black;')
        scroll_layout.addWidget(self.email_input)

        # Detailed Mode Options Container
        self.detailed_options_container = QWidget()
        detailed_options_layout = QVBoxLayout(self.detailed_options_container)
        detailed_options_layout.setContentsMargins(0,0,0,0)

        # Literature Group
        self.lit_group = QGroupBox('Literature Screening Configuration')
        lit_layout = QVBoxLayout()
        year_layout = QHBoxLayout()
        year_layout.addWidget(QLabel('Pub Year Range:'))
        self.start_year_input = QSpinBox(); self.start_year_input.setRange(1900, 2100); self.start_year_input.setValue(2015); self.start_year_input.setSpecialValueText('No Limit')
        self.end_year_input = QSpinBox(); self.end_year_input.setRange(1900, 2100); self.end_year_input.setValue(datetime.now().year)
        year_layout.addWidget(self.start_year_input); year_layout.addWidget(QLabel('to')); year_layout.addWidget(self.end_year_input)
        lit_layout.addLayout(year_layout)
        lit_layout.addWidget(QLabel('Literature Limit per Protein *'))
        self.lit_limit_input = QSpinBox(); self.lit_limit_input.setRange(1, 100); self.lit_limit_input.setValue(15)
        lit_layout.addWidget(self.lit_limit_input)
        self.prioritize_reviews_check = QCheckBox('Prioritize Review Articles'); self.prioritize_reviews_check.setChecked(True)
        lit_layout.addWidget(self.prioritize_reviews_check)
        self.lit_group.setLayout(lit_layout)
        detailed_options_layout.addWidget(self.lit_group)

        # Literature Quality Group
        self.lit_quality_group = QGroupBox('Literature Quality Control (Europe PMC)')
        lit_quality_layout = QVBoxLayout()
        self.open_access_only_check = QCheckBox('Open access only'); self.open_access_only_check.setChecked(False)
        lit_quality_layout.addWidget(self.open_access_only_check)
        self.fulltext_priority_check = QCheckBox('Prioritize articles with full text'); self.fulltext_priority_check.setChecked(True)
        lit_quality_layout.addWidget(self.fulltext_priority_check)
        min_citation_layout = QHBoxLayout()
        min_citation_layout.addWidget(QLabel('Minimum Citations:'))
        self.min_citation_spinbox = QSpinBox(); self.min_citation_spinbox.setRange(0, 1000); self.min_citation_spinbox.setValue(0); self.min_citation_spinbox.setSpecialValueText('No Limit')
        min_citation_layout.addWidget(self.min_citation_spinbox); min_citation_layout.addStretch()
        lit_quality_layout.addLayout(min_citation_layout)
        self.lit_quality_group.setLayout(lit_quality_layout)
        detailed_options_layout.addWidget(self.lit_quality_group)

        # Pollutant Background Group
        self.pollutant_lit_group = QGroupBox('Pollutant Background Literature')
        pollutant_lit_layout = QVBoxLayout()
        self.include_pollutant_bg_check = QCheckBox('Include pollutant background search'); self.include_pollutant_bg_check.setChecked(True)
        self.include_pollutant_bg_check.setToolTip('Searches toxicology reviews, epidemiology for better context.')
        pollutant_lit_layout.addWidget(self.include_pollutant_bg_check)
        pollutant_limit_layout = QHBoxLayout()
        pollutant_limit_layout.addWidget(QLabel('  Pollutant literature limit:'))
        self.pollutant_lit_limit_spinbox = QSpinBox(); self.pollutant_lit_limit_spinbox.setRange(5, 50); self.pollutant_lit_limit_spinbox.setValue(10); self.pollutant_lit_limit_spinbox.setSuffix(' articles')
        pollutant_limit_layout.addWidget(self.pollutant_lit_limit_spinbox); pollutant_limit_layout.addStretch()
        pollutant_lit_layout.addLayout(pollutant_limit_layout)
        self.pollutant_lit_group.setLayout(pollutant_lit_layout)
        detailed_options_layout.addWidget(self.pollutant_lit_group)

        # Writing Mode Group
        self.writing_mode_group = QGroupBox('Mini-Review Writing Mode')
        writing_layout = QVBoxLayout()
        self.standard_writing_radio = QRadioButton('Standard Writing (~2-3 min)'); self.standard_writing_radio.setChecked(True)
        writing_layout.addWidget(self.standard_writing_radio)
        writing_layout.addWidget(QLabel('  * Basic synthesis, standard review quality'))
        self.deep_writing_radio = QRadioButton('Deep Research Writing (~5-8 min)')
        writing_layout.addWidget(self.deep_writing_radio)
        writing_layout.addWidget(QLabel('  * Deeper analysis, gap ID, iterative writing'))
        self.writing_mode_group.setLayout(writing_layout)
        detailed_options_layout.addWidget(self.writing_mode_group)

        # Experimental Data Group
        self.exp_data_group = QGroupBox('Experimental Data (optional)')
        exp_data_layout = QVBoxLayout()
        exp_data_layout.addWidget(QLabel('Add preliminary data (WB, qPCR etc.) to improve AI accuracy.'))
        exp_button_layout = QHBoxLayout()
        self.exp_data_btn = QPushButton('Add/Edit Experimental Data')
        self.exp_data_btn.clicked.connect(self.open_experimental_data_dialog)
        exp_button_layout.addWidget(self.exp_data_btn)
        self.exp_data_status = QLabel('No data added')
        self.exp_data_status.setStyleSheet('color: gray; font-size: 9pt;')
        exp_button_layout.addWidget(self.exp_data_status); exp_button_layout.addStretch()
        exp_data_layout.addLayout(exp_button_layout)
        self.exp_data_group.setLayout(exp_data_layout)
        #detailed_options_layout.addWidget(self.exp_data_group)
        self.exp_data_group.setVisible(False)
        # Expert Panel Group
        self.expert_panel_group = QGroupBox('Expert Review Panel')
        expert_panel_layout = QVBoxLayout()
        info_label = QLabel(
            '<b style="color: #1976D2;">Required experts (must select):</b> '
            'PI, Molecular Toxicologist, Scientific Critic<br>'
            '<b>Optional experts (select as needed):</b> Other specialists'
        )
        info_label.setWordWrap(True)
        info_label.setStyleSheet('padding: 5px; background-color: #E3F2FD; border-radius: 3px; margin-bottom: 5px;')
        expert_panel_layout.addWidget(info_label)
        self.expert_checks = {}
        expert_descriptions = {
            'PI (Principal Investigator)': 'PI - Leads research, guides discussion, synthesizes findings',
            'Molecular Toxicologist': 'Molecular Toxicologist - Xenobiotic metabolism, cellular damage mechanisms',
            'Scientific Critic': 'Scientific Critic - Challenges assumptions, ensures evidence-based analysis',
            'Computational Biologist': 'Computational Biologist - Pathway analysis, bioinformatics, networks',
            'Cell Biologist': 'Cell Biologist - Cellular signaling, stress responses, cell death',
            'Pharmacologist': 'Pharmacologist - Drug-target interactions, dose-response relationships',
            'Epidemiologist': 'Epidemiologist - Population studies, exposure assessment, risk analysis',
            'Pathologist': 'Pathologist - Tissue pathology, disease mechanisms, biomarkers',
            'Developmental Toxicologist': 'Developmental Toxicologist - Developmental effects, teratology',
            'Immunotoxicologist': 'Immunotoxicologist - Immune system effects, inflammation',
            'Biochemist': 'Biochemist - Protein structure-function, enzyme kinetics'
        }
        required_experts = {
            'PI (Principal Investigator)',
            'Molecular Toxicologist',
            'Scientific Critic'
        }
        for expert_key, expert_desc in expert_descriptions.items():
            check = QCheckBox(expert_desc)
            if expert_key in required_experts:
                check.setChecked(True)
                check.setEnabled(False)
                check.setStyleSheet('font-weight: bold; color: #1565C0;')
            else:
                check.setChecked(False)
            self.expert_checks[expert_key] = check
            expert_panel_layout.addWidget(check)
        self.expert_panel_group.setLayout(expert_panel_layout)
        detailed_options_layout.addWidget(self.expert_panel_group)

        scroll_layout.addWidget(self.detailed_options_container)

        # AI Configuration
        scroll_layout.addWidget(QLabel('AI Analysis Model'))
        self.model_combo = QComboBox(); self.model_combo.addItems(['Gemini 2.5 Flash', 'Gemini 2.5 Pro', 'qwen-plus', 'qwen-max', 'deepseek-chat', 'deepseek-reasoner']); self.model_combo.setCurrentText('Gemini 2.5 Flash')
        scroll_layout.addWidget(self.model_combo)

        scroll_layout.addWidget(QLabel('API Key *'))
        api_layout = QHBoxLayout()
        self.api_key_input = QLineEdit(); self.api_key_input.setEchoMode(QLineEdit.Password); self.api_key_input.setStyleSheet('background-color: white; color: black;')
        api_layout.addWidget(self.api_key_input)
        self.show_api_btn = QPushButton('Show'); self.show_api_btn.setCheckable(True); self.show_api_btn.setMaximumWidth(60); self.show_api_btn.toggled.connect(self.toggle_api_visibility)
        api_layout.addWidget(self.show_api_btn)
        scroll_layout.addLayout(api_layout)

        scroll_layout.addWidget(QLabel('Proxy URL (optional)'))
        self.proxy_input = QLineEdit(); self.proxy_input.setPlaceholderText('e.g., http://127.0.0.1:7890'); self.proxy_input.setStyleSheet('background-color: white; color: black;')
        scroll_layout.addWidget(self.proxy_input)

        test_btn = QPushButton('Test API Connection')
        test_btn.clicked.connect(self.test_api_connection)
        scroll_layout.addWidget(test_btn)

        # Final Scroll Area Setup
        scroll_layout.addStretch()
        scroll.setWidget(scroll_widget)
        layout.addWidget(scroll)

        # Start/Stop Buttons
        button_layout = QHBoxLayout()
        self.start_btn = QPushButton('Start Screening')
        self.start_btn.setStyleSheet(
            'background-color: #4CAF50; color: white; padding: 12px 30px; '
            'font-weight: bold; font-size: 14pt;'
        )
        self.start_btn.clicked.connect(self.start_analysis)
        button_layout.addWidget(self.start_btn)

        self.stop_btn = QPushButton('Stop Analysis')
        self.stop_btn.setStyleSheet(
            'background-color: #f44336; color: white; padding: 12px 30px; '
            'font-weight: bold; font-size: 14pt;'
        )
        self.stop_btn.setEnabled(False)
        self.stop_btn.clicked.connect(self.stop_analysis)
        button_layout.addWidget(self.stop_btn)

        layout.addLayout(button_layout)

        return panel

    def create_right_panel(self) -> QWidget:
        panel = QWidget()
        panel.setStyleSheet('background-color: white;')
        layout = QVBoxLayout(panel)

        title = QLabel('Analysis Progress & Results')
        title.setStyleSheet('font-size: 16pt; font-weight: bold; padding: 10px; color: #333;')
        layout.addWidget(title)

        self.progress_widget = QFrame()
        self.progress_widget.setStyleSheet(
            'background-color: #f5f5f5; padding: 15px; border-radius: 8px; '
            'border: 1px solid #ddd;'
        )
        self.progress_widget.setVisible(False)
        progress_layout = QVBoxLayout(self.progress_widget)

        progress_header = QHBoxLayout()
        progress_label = QLabel('Progress:')
        progress_label.setStyleSheet('font-weight: bold; font-size: 12pt; color: #333;')
        progress_header.addWidget(progress_label)
        self.progress_percentage = QLabel('0%')
        self.progress_percentage.setStyleSheet('font-weight: bold; font-size: 12pt; color: #4CAF50;')
        progress_header.addWidget(self.progress_percentage)
        progress_header.addStretch()
        progress_layout.addLayout(progress_header)

        self.progress_bar = QProgressBar()
        self.progress_bar.setStyleSheet(
            'QProgressBar {border: 2px solid #ddd; border-radius: 5px; text-align: center; height: 25px;}'
            'QProgressBar::chunk {background-color: #4CAF50; border-radius: 3px;}'
        )
        progress_layout.addWidget(self.progress_bar)

        self.progress_text = QLabel('Initializing...')
        self.progress_text.setStyleSheet('font-size: 11pt; color: #333; font-weight: bold;')
        progress_layout.addWidget(self.progress_text)

        self.progress_detail = QLabel('')
        self.progress_detail.setStyleSheet('font-size: 9pt; color: #666; font-style: italic;')
        self.progress_detail.setWordWrap(True)
        progress_layout.addWidget(self.progress_detail)

        layout.addWidget(self.progress_widget)

        log_header = QHBoxLayout()
        log_label = QLabel('Execution Log:')
        log_label.setStyleSheet('color: #333;')
        log_header.addWidget(log_label)
        log_header.addStretch()

        clear_log_btn = QPushButton('Clear Log')
        clear_log_btn.clicked.connect(self.clear_log)
        log_header.addWidget(clear_log_btn)

        self.view_notebook_btn = QPushButton('View Research Notebook')
        self.view_notebook_btn.setStyleSheet('background-color: #2196F3; color: white; font-weight: bold;')
        self.view_notebook_btn.setVisible(False)
        self.view_notebook_btn.clicked.connect(self.view_notebook)
        log_header.addWidget(self.view_notebook_btn)

        self.view_conclusion_btn = QPushButton('View Final Conclusions')
        self.view_conclusion_btn.setStyleSheet('background-color: #9C27B0; color: white; font-weight: bold;')
        self.view_conclusion_btn.setVisible(False)
        self.view_conclusion_btn.clicked.connect(self.view_conclusion)
        log_header.addWidget(self.view_conclusion_btn)

        self.export_db_csv_btn = QPushButton('Export Database CSV')
        self.export_db_csv_btn.setStyleSheet('background-color: #00897B; color: white; font-weight: bold;') # Teal
        self.export_db_csv_btn.setVisible(False)
        self.export_db_csv_btn.clicked.connect(self.export_database_csv)
        log_header.addWidget(self.export_db_csv_btn)

        self.view_mini_review_btn = QPushButton('View Mini-Review')
        self.view_mini_review_btn.setStyleSheet('background-color: #3F51B5; color: white; font-weight: bold;') # Indigo
        self.view_mini_review_btn.setVisible(False)
        self.view_mini_review_btn.clicked.connect(self.view_mini_review) # New slot
        log_header.addWidget(self.view_mini_review_btn)

        layout.addLayout(log_header)

        self.log_text = QTextEdit()
        self.log_text.setReadOnly(True)
        self.log_text.setStyleSheet(
            'background-color: white; '
            'color: #333333; '
            'font-family: Consolas, monospace; '
            'font-size: 9pt; '
            'border: 1px solid #ddd;'
        )
        layout.addWidget(self.log_text)

        return panel

    def parse_proteins(self, text: str) -> List[str]:
        """Parse protein names from text"""
        text = text.replace(',', '\n').replace(';', '\n')
        proteins = [line.strip() for line in text.split('\n') if line.strip()]
        return proteins

    def update_protein_count(self):
        """Update protein count label and mode hint"""
        proteins = self.parse_proteins(self.proteins_input.toPlainText())
        count = len(proteins)
        self.protein_count_label.setText(f'{count} proteins')

        is_quick = self.quick_mode_radio.isChecked()
        if is_quick:
            if count > 0 and count < 20:
                self.mode_hint_label.setText('! Consider Detailed Screening for small sets (<20 proteins)')
            else:
                self.mode_hint_label.setText('')
        else:
            if count > 15:
                self.mode_hint_label.setText('! Large set detected. Detailed mode may take 30+ min. Consider Quick mode first.')
            else:
                self.mode_hint_label.setText('')

    def on_species_changed(self):
        """Update species data availability hint"""
        species = self.species_combo.currentText()
        hints = {
            'Human (Homo sapiens)': 'Best data availability: KEGG, UniProt, GTEx, HPA',
            'Mouse (Mus musculus)': 'Good data: KEGG, UniProt (Expression data via Human Ortholog)',
            'Zebrafish (Danio rerio)': 'Limited data: KEGG, UniProt (Expression data via Human Ortholog)'
        }
        self.species_data_hint.setText(hints.get(species, ''))

    def on_mode_changed(self):
        """Update UI visibility based on selected mode."""
        is_quick = self.quick_mode_radio.isChecked()

        self.quick_options_widget.setVisible(is_quick)
        self.detailed_options_container.setVisible(not is_quick)
        self.update_protein_count()

    def open_experimental_data_dialog(self):
        dialog = ExperimentalDataDialog(self.experimental_data, self)

        if dialog.exec_() == QDialog.Accepted:
            self.experimental_data = dialog.get_data()

            if self.experimental_data:
                self.exp_data_status.setText(f'Data added ({len(self.experimental_data)} chars)')
                self.exp_data_status.setStyleSheet('color: green; font-size: 9pt; font-weight: bold;')
            else:
                self.exp_data_status.setText('No data added')
                self.exp_data_status.setStyleSheet('color: gray; font-size: 9pt;')

    def toggle_api_visibility(self, checked):
        if checked:
            self.api_key_input.setEchoMode(QLineEdit.Normal)
            self.show_api_btn.setText('Hide')
        else:
            self.api_key_input.setEchoMode(QLineEdit.Password)
            self.show_api_btn.setText('Show')

    def start_analysis(self):
        pollutant = self.pollutant_input.text().strip()
        proteins_text = self.proteins_input.toPlainText().strip()
        tissue = self.tissue_combo.currentText()
        species = self.species_combo.currentText()
        api_key = self.api_key_input.text().strip()
        is_quick_mode = self.quick_mode_radio.isChecked()
        email = self.email_input.text().strip()

        required_fields = {
            'Pollutant': pollutant,
            'Candidate Proteins': proteins_text,
            'Target Tissue': tissue if tissue != 'Select Tissue' else '',
            'API Key': api_key
        }
        missing = [name for name, value in required_fields.items() if not value]
        if missing:
            QMessageBox.warning(self, 'Input Error', f'Please fill in all required fields: {", ".join(missing)}.')
            return

        if not email or "@" not in email:
            QMessageBox.warning(self, 'Input Error', 'Please enter a valid email address. It is required for Detailed mode.')
            return

        proteins = self.parse_proteins(proteins_text)
        if not proteins:
            QMessageBox.warning(self, 'Input Error', 'Please enter at least one valid protein name.')
            return

        phenotypes_text = self.phenotypes_input.toPlainText().strip()
        phenotypes = [p.strip() for p in phenotypes_text.split('\n') if p.strip()]

        selected_experts = []
        if not is_quick_mode:
            selected_experts = [k for k, v in self.expert_checks.items() if v.isChecked()]
            required_experts = {'PI (Principal Investigator)', 'Molecular Toxicologist', 'Scientific Critic'}
            for req in required_experts:
                 if req not in selected_experts: selected_experts.append(req)

            if len(selected_experts) < 3:
                QMessageBox.warning(self, 'Expert Selection', 'At least 3 experts (including required) must be selected for detailed review.')
                return

        display_model = self.model_combo.currentText()
        model_mapping = {
            'Gemini 2.5 Flash': 'gemini-2.5-flash-preview-09-2025',
            'Gemini 2.5 Pro': 'gemini-2.5-pro-preview-06-05',
            'qwen-plus': 'qwen-plus',
            'qwen-max': 'qwen-max',
            'deepseek-chat': 'deepseek-chat',
            'deepseek-reasoner': 'deepseek-reasoner'
        }
        model_name = model_mapping.get(display_model, display_model)

        self.log_buffer = []
        self.log_text.clear()
        self.progress_widget.setVisible(True)
        self.progress_bar.setValue(0)
        self.progress_percentage.setText('0%')
        self.progress_text.setText('Initializing...')
        self.progress_detail.setText('')
        self.view_notebook_btn.setVisible(False)
        self.view_conclusion_btn.setVisible(False)
        self.export_db_csv_btn.setVisible(False)
        self.view_mini_review_btn.setVisible(False)

        params = {
            'pollutant': pollutant,
            'proteins': proteins,
            'target_phenotypes': phenotypes,
            'target_tissue': tissue,
            'species': species,
            'email': email,
            'model': model_name,
            'api_key': api_key,
            'proxy': self.proxy_input.text().strip() or None,
            'is_quick_mode': is_quick_mode,
            'species_taxon': {
                'Human (Homo sapiens)': 9606,
                'Mouse (Mus musculus)': 10090,
                'Zebrafish (Danio rerio)': 7955
            }.get(species, 9606),
            'species_kegg': {
                'Human (Homo sapiens)': 'hsa',
                'Mouse (Mus musculus)': 'mmu',
                'Zebrafish (Danio rerio)': 'dre'
            }.get(species, 'hsa')
        }

        if is_quick_mode:
            params['top_n'] = self.top_n_spinbox.value()
            params['consistency_rounds'] = self.consistency_spinbox.value()
        else:
            params.update({
                'start_year': self.start_year_input.value() if self.start_year_input.value() > 1900 else '',
                'end_year': self.end_year_input.value() if self.end_year_input.value() > 1900 else '',
                'lit_limit': self.lit_limit_input.value(),
                'prioritize_reviews': self.prioritize_reviews_check.isChecked(),
                'open_access_only': self.open_access_only_check.isChecked(),
                'fulltext_priority': self.fulltext_priority_check.isChecked(),
                'min_citations': self.min_citation_spinbox.value(),
                'include_pollutant_bg': self.include_pollutant_bg_check.isChecked(),
                'pollutant_lit_limit': self.pollutant_lit_limit_spinbox.value(),
                'deep_writing_mode': self.deep_writing_radio.isChecked(),
                'experimental_data': self.experimental_data,
                'selected_experts': selected_experts
            })

        self.session_data = {
            'inputs': params,
            'results': None,
            'timestamp': datetime.now().isoformat()
        }

        logging.info(f"Starting analysis: {params['pollutant']} | {len(params['proteins'])} proteins | Mode: {'Quick' if is_quick_mode else 'Detailed'}")
        if is_quick_mode:
            logging.info(f"Quick Mode Options: Top N={params['top_n']}, Consistency Rounds={params['consistency_rounds']}")

        method_to_run = 'quick_screening' if is_quick_mode else 'detailed_screening'
        self.analysis_thread = AnalysisThread(self.engine, method_to_run, params)
        self.analysis_thread.start()

        self.start_btn.setEnabled(False)
        self.stop_btn.setEnabled(True)

    def stop_analysis(self):
        if self.analysis_thread and self.analysis_thread.isRunning():
            reply = QMessageBox.question(
                self, 'Confirm Stop',
                'Are you sure you want to stop the analysis?',
                QMessageBox.Yes | QMessageBox.No
            )
            if reply == QMessageBox.Yes:
                if hasattr(self.engine, 'stop'):
                     self.engine.stop()
                if hasattr(self.analysis_thread, 'stop'):
                     self.analysis_thread.stop()
                self.add_log('User requested analysis stop.', 'warning')
                self.start_btn.setEnabled(True)
                self.stop_btn.setEnabled(False)
                self.progress_text.setText('Analysis Stopped by User')
                self.progress_detail.setText('')


    def test_api_connection(self):
        api_key = self.api_key_input.text().strip()
        proxy = self.proxy_input.text().strip() or None
        display_model = self.model_combo.currentText()

        model_mapping = {
            'Gemini 2.5 Flash': 'gemini-2.5-flash-preview-09-2025',
            'Gemini 2.5 Pro': 'gemini-2.5-pro-preview-06-05',
            'qwen-plus': 'qwen-plus',
            'qwen-max': 'qwen-max',
            'deepseek-chat': 'deepseek-chat',
            'deepseek-reasoner': 'deepseek-reasoner'
        }
        model_name = model_mapping.get(display_model, display_model)

        if not api_key:
            QMessageBox.warning(self, 'Missing API Key', 'Please enter your API key first.')
            return

        self.add_log(f'Testing API connection for model: {model_name}...', 'info')

        self.test_worker = ApiConnectionTestWorker(proxy, model_name, api_key)
        self.test_thread = QThread()
        self.test_worker.moveToThread(self.test_thread)

        self.test_worker.finished.connect(self.on_test_finished)
        self.test_worker.finished.connect(self.test_thread.quit)
        self.test_worker.finished.connect(self.test_worker.deleteLater)
        self.test_thread.finished.connect(self.test_thread.deleteLater)

        self.test_thread.started.connect(self.test_worker.run)
        self.test_thread.start()

    def on_test_finished(self, success: bool, message: str):
        if success:
            QMessageBox.information(self, 'Test Successful', f'API connection successful!\n\n{message}')
            self.add_log(f'API connection test PASSED: {message}', 'success')
        else:
            QMessageBox.critical(self, 'Test Failed', f'API connection failed.\n\n{message}')
            self.add_log(f'API connection test FAILED: {message}', 'error')

    @pyqtSlot(str, str)
    def add_log(self, message: str, level: str = 'info'):
        colors = {
            'info': '#555555',
            'success': '#4CAF50',
            'warning': '#FF9800',
            'error': '#f44336',
            'debug': '#9E9E9E'
        }
        color = colors.get(level, '#555555')

        timestamp = datetime.now().strftime('%H:%M:%S')

        escaped_message = html.escape(message)
        formatted_message = escaped_message.replace('\n', '<br>')

        formatted_html = f'<span style="color: #888;">[{timestamp}]</span> <span style="color: {color};">{formatted_message}</span><br>'

        self.log_buffer.append(formatted_html)
        if len(self.log_buffer) > 1000:
            self.log_buffer.pop(0)

        self.log_text.append(formatted_html)
        self.log_text.verticalScrollBar().setValue(self.log_text.verticalScrollBar().maximum())

    @pyqtSlot(str, str)
    def add_detailed_log(self, stage: str, content: str):
        self.add_log(f'[{stage}]\n{content}', 'info')


    @pyqtSlot(int, str, str)
    def update_progress(self, value: int, text: str, detail: str = ''):
        self.progress_bar.setValue(value)
        self.progress_percentage.setText(f'{value}%')
        self.progress_text.setText(text)
        self.progress_detail.setText(detail)

    def clear_log(self):
        self.log_text.clear()
        self.log_buffer = []

    @pyqtSlot(dict)
    def on_quick_screening_complete(self, session_data: dict):
        self.add_log('Quick screening complete. Displaying top results...', 'success')
        self.session_data = session_data

        screening_result = session_data.get('quick_screening_result', {})
        dialog = QuickScreeningResultDialog(screening_result, session_data.get('inputs', {}), self)
        result = dialog.exec_()

        if result == QDialog.Accepted:
            selected_proteins = dialog.get_selected_proteins()
            if not selected_proteins:
                QMessageBox.warning(self, 'No Selection', 'No proteins selected. Completing quick mode.')
                self.on_analysis_complete(self.session_data)
                return

            self.add_log(f'User selected {len(selected_proteins)} proteins for detailed screening.', 'info')

            proteins_text = '\n'.join(selected_proteins)
            self.proteins_input.setPlainText(proteins_text)

            self.detailed_mode_radio.setChecked(True)

            self.session_data['inputs']['proteins'] = selected_proteins

            self.progress_widget.setVisible(False)
            self.start_btn.setEnabled(True)
            self.stop_btn.setEnabled(False)

            self.add_log(
                f'Selected {len(selected_proteins)} proteins have been loaded into Detailed Screening input.\n'
                'Please review/adjust parameters and click START to begin detailed analysis.',
                'info'
            )

            QMessageBox.information(
                self,
                'Ready for Detailed Screening',
                f'{len(selected_proteins)} proteins have been loaded into the input box.\n\n'
                'The mode has been switched to Detailed Screening.\n\n'
                'Please review the parameters (literature count, expert panel settings, etc.) '
                'and click the START button when ready.'
            )

        else:
            self.add_log('User skipped detailed screening. Analysis complete.', 'info')
            self.on_analysis_complete(self.session_data)


    @pyqtSlot(dict)
    def show_literature_review(self, literature_data: dict):
        self.add_log('Literature search complete. Review retrieved articles...', 'success')

        dialog = LiteratureReviewDialog(literature_data, self.session_data['inputs'], self)
        result = dialog.exec_()

        if result == QDialog.Accepted:
            selected_lit = dialog.get_selected_literature()
            self.add_log(f'User confirmed literature selection for {len(selected_lit)} proteins.', 'info')
            self.start_btn.setEnabled(False)
            self.stop_btn.setEnabled(True)
            self.analysis_thread = AnalysisThread(self.engine, 'continue_detailed_screening', selected_lit)
            self.analysis_thread.start()
        else:
            self.add_log('Literature review cancelled by user.', 'warning')
            if self.engine:
                 self.engine.stop()
            self.on_error('Literature review cancelled.')


    @pyqtSlot(dict)
    def on_analysis_complete(self, session_data: dict):
        self.session_data = session_data
        self.add_log('Analysis complete!', 'success')

        self.start_btn.setEnabled(True)
        self.stop_btn.setEnabled(False)

        self.progress_text.setText('Analysis Complete')
        self.progress_detail.setText('')

        results = session_data.get('results', {})
        
        # --- Notebook Button ---
        notebook_obj = session_data.get('notebook')
        if notebook_obj and hasattr(notebook_obj, 'notebook_data'):
            self.view_notebook_btn.setVisible(True)

        # --- Database CSV Button ---
        if session_data.get('database_csv_path'):
            self.export_db_csv_btn.setVisible(True)

        # --- Concise Report Button ---
        if session_data.get('final_concise_report'):
            self.view_conclusion_btn.setVisible(True)
            self.view_conclusion_btn.setText('View Concise Report')
        
        # --- MODIFICATION START (Using new 'Mini-Review' key and logic) ---
        # --- Mini-Review Button (Enhanced logic) ---
        final_eval = session_data.get('final_evaluation', {})
        
        # Check for full (deep) mini-review
        has_full_review = final_eval.get('full_review_text')
        
        # Check for standard review data (or deep review error structure)
        mini_review_data = final_eval.get('Mini-Review', {})  # <-- Use new key
        has_standard_review_or_error = (
            isinstance(mini_review_data, dict) and 
            len(mini_review_data) > 0
        )
        
        # Check if it's a valid review (not *just* an error)
        has_valid_review = (
            has_standard_review_or_error and 
            'error' not in mini_review_data 
        )
        
        if has_full_review or has_valid_review:
            self.view_mini_review_btn.setVisible(True)
            self.add_log(f'Mini-review available (Full: {bool(has_full_review)}, Standard: {bool(has_valid_review)})', 'info')
        else:
            # If data exists but contains an error, log it
            if final_eval.get('Mini-Review', {}).get('error'): # <-- Use new key
                self.add_log(f'Mini-review generation failed: {final_eval["Mini-Review"]["error"]}', 'warn')
            elif final_eval.get('error'): # Fallback for old deep error structure
                 self.add_log(f'Mini-review generation failed: {final_eval.get("error")}', 'warn')
            else:
                self.add_log('Mini-review data not available in final_evaluation', 'warn')
        # --- MODIFICATION END ---

        try:
            session_id = f"session_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
            self.current_session_id = session_id
            self.session_manager.save_session(session_data, session_id)
            self.add_log(f'Session saved: {session_id}', 'info')
        except Exception as e:
            logging.error(f"Failed to save session: {e}", exc_info=True)
            self.add_log(f"Failed to save session: {e}", 'error')

        QMessageBox.information(
            self, 'Analysis Complete',
            'Screening analysis finished successfully!\n\n'
            'Use the export buttons to save results.'
        )

    @pyqtSlot(str)
    def on_error(self, error_message: str):
        self.add_log(f'ERROR: {error_message}', 'error')

        self.start_btn.setEnabled(True)
        self.stop_btn.setEnabled(False)

        self.progress_text.setText('Analysis Failed')
        self.progress_detail.setText(error_message)

        if self.session_data and self.session_data.get('database_csv_path'):
            self.add_log('ERROR: Analysis failed, but database CSV file was generated.', 'warning')
            self.export_db_csv_btn.setVisible(True)

        if self.session_data and 'inputs' in self.session_data:
             try:
                 session_id = f"error_session_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
                 self.session_manager.save_session(self.session_data, session_id)
                 self.add_log(f'Partial session data saved due to error: {session_id}', 'warning')
             except Exception as e:
                 logging.error(f"Failed to save partial session data: {e}", exc_info=True)
                 self.add_log(f"Failed to save partial session data: {e}", 'error')

        QMessageBox.critical(
            self, 'Analysis Error',
            f'An error occurred during analysis:\n\n{error_message}\n\n'
            f'Check the log for details. Partial session data may have been saved.'
        )

    def export_database_csv(self):
        """Allows user to copy the generated database CSV to a location"""
        if not self.session_data or not self.session_data.get('database_csv_path'):
            QMessageBox.warning(self, 'No Data', 'No database CSV file path found in session.')
            return

        source_path = self.session_data['database_csv_path']
        
        if not os.path.exists(source_path):
            QMessageBox.critical(self, 'File Not Found', f'Could not find the generated CSV file:\n{source_path}\n\nIt may have been moved or deleted.')
            return

        # Prepare default filename
        inputs_data = self.session_data.get('inputs', {})
        safe_pollutant = sanitize_filename(inputs_data.get('pollutant', 'Analysis'))
        timestamp = datetime.now().strftime('%Y%m%d')
        default_filename = f"Database_Export_{safe_pollutant}_{timestamp}.csv"

        filename, _ = QFileDialog.getSaveFileName(
            self, 'Save Database CSV As', default_filename, 'CSV Files (*.csv)'
        )

        if filename:
            try:
                shutil.copy(source_path, filename)
                QMessageBox.information(self, 'Success', f'Database CSV exported successfully to:\n{filename}')
            except Exception as e:
                logging.error(f"Failed to copy CSV file: {e}", exc_info=True)
                QMessageBox.critical(self, 'Error', f'File copy failed:\n{str(e)}')

    def view_notebook(self):
        if not self.session_data or 'notebook' not in self.session_data:
             QMessageBox.warning(self, 'No Data', 'No session results or notebook available.')
             return

        notebook_obj = self.session_data.get('notebook')

        if not notebook_obj or not hasattr(notebook_obj, 'notebook_data'):
             QMessageBox.warning(self, 'No Data', 'Research notebook data not found or invalid.')
             return
        
        try:
             md_content = self._get_notebook_markdown(notebook_obj)
             
             safe_pollutant = sanitize_filename(self.session_data.get('inputs', {}).get('pollutant', 'notebook'))
             default_filename = f"Research_Notebook_{safe_pollutant}.md"
             
             dialog = SaveableTextDialog(
                 title='Research Notebook',
                 content=md_content,
                 default_filename=default_filename,
                 is_markdown=True,
                 parent=self
             )
             dialog.exec_()
             
        except Exception as e:
             QMessageBox.critical(self, "Error", f"Error generating notebook view:\n{e}")


    def _get_notebook_markdown(self, notebook_obj) -> str:
        """Helper to get notebook content as Markdown string"""
        import tempfile
        import io
        temp_path = None
        try:
            fd, temp_path = tempfile.mkstemp(suffix=".md")
            os.close(fd)

            notebook_obj.export_to_markdown(temp_path)

            with open(temp_path, 'r', encoding='utf-8') as f:
                content = f.read()
            return content
        except Exception as e:
            logging.error(f"Failed to get notebook markdown content: {e}")
            return f"Error creating notebook preview: {e}"
        finally:
            if temp_path and os.path.exists(temp_path):
                try:
                    os.remove(temp_path)
                except OSError as rm_e:
                     logging.warning(f"Could not remove temporary notebook file '{temp_path}': {rm_e}")


    def view_conclusion(self):
        """Displays the new concise report."""
        if not self.session_data:
            QMessageBox.warning(self, 'No Data', 'No session data available.')
            return

        display_content = self.session_data.get('final_concise_report')

        if not display_content:
            QMessageBox.warning(self, 'No Data', 'Concise report is not available. (final_concise_report is missing)')
            return

        safe_pollutant = sanitize_filename(self.session_data.get('inputs', {}).get('pollutant', 'report'))
        default_filename = f"Concise_Report_{safe_pollutant}.md"

        dialog = SaveableTextDialog(
            title='Final Concise Report',
            content=display_content,
            default_filename=default_filename,
            is_markdown=True,
            parent=self
        )
        dialog.exec_()

    def view_mini_review(self):
        """Displays the full, detailed Mini-Review."""
        if not self.session_data:
            QMessageBox.warning(self, 'No Data', 'No session data available.')
            return

        final_eval = self.session_data.get('final_evaluation', {})
        
        display_content = final_eval.get('full_review_text') 
        
        if not display_content:
            # --- MODIFICATION START (Using new 'Mini-Review' key) ---
            mini_review_dict = final_eval.get('Mini-Review', {}) # <-- Use new key
            
            if mini_review_dict and not mini_review_dict.get('error'):
                parts = [
                    f"# Mini-Review (Standard Mode)\n",
                    f"## Executive Summary\n{mini_review_dict.get('executive_summary', 'N/A')}",
                    f"\n## Introduction\n{mini_review_dict.get('introduction', 'N/A')}",
                    "\n## Top Protein Analysis",
                ]
                for item in mini_review_dict.get('top_protein_analysis', []):
                    parts.append(f"### {item.get('protein', 'N/A')} (Rank {item.get('rank', '?')})\n{item.get('analysis_text', 'N/A')}")
                
                parts.append(f"\n## Integrated Mechanism\n{mini_review_dict.get('integrated_mechanism', 'N/A')}")
                parts.append(f"\n## Conclusion & Future Directions\n{mini_review_dict.get('conclusion_future_directions', 'N/A')}")
                parts.append(f"\n## Research Plans\n")
                for item in mini_review_dict.get('research_plans', []):
                    parts.append(f"### {item.get('protein', 'N/A')}\n{item.get('plan', 'N/A')}")
                    
                display_content = "\n".join(parts)
            else:
                display_content = final_eval.get('Mini-Review', {}).get('error', 'Mini-Review data is not available or generation failed.') # <-- Use new key
            # --- MODIFICATION END ---

        if not display_content:
            QMessageBox.warning(self, 'No Data', 'Mini-Review data not found in session.')
            return

        safe_pollutant = sanitize_filename(self.session_data.get('inputs', {}).get('pollutant', 'review'))
        default_filename = f"Full_Mini_Review_{safe_pollutant}.md"

        dialog = SaveableTextDialog(
            title='Full Mini-Review',
            content=display_content,
            default_filename=default_filename,
            is_markdown=True,
            parent=self
        )
        dialog.exec_()


if __name__ == '__main__':
    log_dir = Path('logs'); log_dir.mkdir(exist_ok=True)
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    log_file = log_dir / f'screening_app_{timestamp}.log'
    try:
        setup_logging(log_file=str(log_file), level=logging.INFO)
        logging.info("Application starting...")
    except Exception as e:
         print(f"FATAL: Could not configure file logging to {log_file}. Error: {e}")
         logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
         logging.error(f"File logging failed, using console only. Error: {e}")

    app = QApplication(sys.argv)
    app.setStyle('Fusion')

    window = TwoStageMainWindow()
    window.show()

    exit_code = app.exec_()
    logging.info(f"Application exiting with code {exit_code}.")
    sys.exit(exit_code)