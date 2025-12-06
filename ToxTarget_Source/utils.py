# -*- coding: utf-8 -*-
"""
Utility Functions Module
- Manages caching, rate limiting, and concurrency.
- Provides data validation and session management.
- Includes ResearchNotebook for recording and exporting findings.
- Includes LiteratureManager for report generation.
- Includes ExpressionAnalyzer for tissue expression data analysis.
- Includes exporters for quick screening reports.
- Includes Token management utilities for AI API safety.
"""

import sqlite3
import time
import json
import asyncio
import pickle
import logging
import os
from typing import Dict, List, Any, Callable, Optional, Tuple
from datetime import datetime
from pathlib import Path
import re
import csv

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Filename Sanitization

def sanitize_filename(filename: str, replacement: str = '_') -> str:
    """
    Sanitize a filename by removing or replacing invalid characters.
    
    Args:
        filename: The filename to sanitize
        replacement: Character to replace invalid characters with (default: '_')
    
    Returns:
        Sanitized filename safe for all operating systems
    """
    invalid_chars = r'[\\/:*?"<>|]'
    sanitized = re.sub(invalid_chars, replacement, filename)
    
    sanitized = sanitized.strip('. ')
    
    sanitized = re.sub(f'{re.escape(replacement)}+', replacement, sanitized)
    
    if not sanitized:
        sanitized = 'unnamed'
    
    return sanitized

# Logging Configuration

def setup_logging(log_file: str = 'protein_screening.log', level=logging.INFO):
    """Configure global logging system"""
    logging.basicConfig(
        level=level,
        format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
        handlers=[
            logging.FileHandler(log_file, encoding='utf-8'),
            logging.StreamHandler()
        ]
    )
    logger = logging.getLogger(__name__)
    logger.info("Logging system initialized")
    return logger


# Token Management Utilities

def estimate_tokens(text: str) -> int:
    """
    Estimate token count for mixed language text
    Rule: 1 token ≈ 4 chars (English), 1 token ≈ 1.5 chars (Chinese)
    """
    if not text:
        return 0
    
    chinese_chars = sum(1 for c in text if '\u4e00' <= c <= '\u9fff')
    other_chars = len(text) - chinese_chars
    
    # Conservative estimation
    estimated_tokens = int(chinese_chars / 1.5 + other_chars / 4)
    
    return estimated_tokens


def truncate_text_by_tokens(
    text: str, 
    max_tokens: int, 
    keep_start: bool = True,
    add_marker: bool = True
) -> str:
    """
    Truncate text to fit within token limit
    
    Args:
        text: Input text
        max_tokens: Maximum allowed tokens
        keep_start: If True, keep beginning; else keep end
        add_marker: If True, add truncation marker
    
    Returns:
        Truncated text
    """
    if not text:
        return text
    
    estimated_tokens = estimate_tokens(text)
    
    if estimated_tokens <= max_tokens:
        return text
    
    # Calculate safe character count (with 90% safety margin)
    ratio = max_tokens / estimated_tokens
    safe_chars = int(len(text) * ratio * 0.9)
    
    if safe_chars <= 0:
        return ""
    
    # Truncate
    if keep_start:
        truncated = text[:safe_chars]
        if add_marker:
            truncated += "\n\n[...Content truncated due to length limit...]"
    else:
        truncated = text[-safe_chars:]
        if add_marker:
            truncated = "[...Earlier content truncated...]\n\n" + truncated
    
    return truncated


def split_text_into_chunks(
    text: str, 
    max_tokens_per_chunk: int, 
    overlap_tokens: int = 200
) -> List[str]:
    """
    Split long text into overlapping chunks
    
    Args:
        text: Input text
        max_tokens_per_chunk: Maximum tokens per chunk
        overlap_tokens: Overlap between chunks (in tokens)
    
    Returns:
        List of text chunks
    """
    if not text:
        return []
    
    total_tokens = estimate_tokens(text)
    
    if total_tokens <= max_tokens_per_chunk:
        return [text]
    
    # Estimate characters per chunk (with 90% safety margin)
    chars_per_chunk = int(len(text) * (max_tokens_per_chunk / total_tokens) * 0.9)
    overlap_chars = int(len(text) * (overlap_tokens / total_tokens) * 0.9)
    
    if chars_per_chunk <= 0:
        return [text[:1000]]  # Emergency fallback
    
    chunks = []
    start = 0
    
    # --- START OF FIX ---
    while start < len(text):
        # Store the start position of this chunk
        prev_start = start 
        
        end = min(start + chars_per_chunk, len(text))
        
        if end < len(text):
            # Look for break points in order of preference
            for separator in ['\n\n', '\n', '. ', '。', '! ', '！', '? ', '？']:
                last_sep = text[start:end].rfind(separator)
                # Accept if found in last 30% of chunk
                if last_sep > chars_per_chunk * 0.7:
                    end = start + last_sep + len(separator)
                    break
        
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)

        # --- FIX: Prevent infinite loop ---
        if end == len(text):
            # We have processed the final chunk
            break 
        
        start = end - overlap_chars
        
        if start <= prev_start:
            # Safety check: If overlap is too large or chunk is too small,
            # 'start' might not advance. Force it to advance to prevent
            # an infinite loop by starting the next chunk at the end of the current one.
            start = end
        # --- END OF FIX ---

    return chunks


def validate_input_length(
    system_prompt: str,
    user_prompt: str,
    max_total_tokens: int,
    logger: Optional[logging.Logger] = None
) -> Tuple[bool, int, str]:
    """
    Validate if combined prompts exceed token limit
    
    Args:
        system_prompt: System prompt text
        user_prompt: User prompt text
        max_total_tokens: Maximum allowed total tokens
        logger: Optional logger for warnings
    
    Returns:
        Tuple of (is_valid, total_tokens, message)
    """
    system_tokens = estimate_tokens(system_prompt)
    user_tokens = estimate_tokens(user_prompt)
    total_tokens = system_tokens + user_tokens
    
    if total_tokens <= max_total_tokens:
        return True, total_tokens, "OK"
    
    message = (
        f"Input exceeds limit: {total_tokens} > {max_total_tokens} tokens. "
        f"System: {system_tokens}, User: {user_tokens}"
    )
    
    if logger:
        logger.warning(message)
    
    return False, total_tokens, message



class PersistentCache:
    """SQLite-based persistent cache"""

    def __init__(self, db_path: str = 'cache.db'):
        self.db_path = db_path
        self._init_db()
        self.logger = logging.getLogger(__name__)

    def _init_db(self):
        """Initialize database schema"""
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS cache (
                key TEXT PRIMARY KEY,
                value TEXT,
                timestamp REAL,
                ttl REAL
            )
        ''')
        conn.commit()
        conn.close()

    def get(self, key: str) -> Optional[Any]:
        """Retrieve cached value by key"""
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        cursor.execute(
            'SELECT value, timestamp, ttl FROM cache WHERE key = ?',
            (key,)
        )
        result = cursor.fetchone()
        conn.close()

        if result:
            value, timestamp, ttl = result
            if time.time() - timestamp < ttl:
                try:
                    return json.loads(value)
                except json.JSONDecodeError:
                    self.logger.warning(f"Failed to parse cached data: {key}")
                    return None
        return None

    def set(self, key: str, value: Any, ttl: float = 86400):
        """Store value in cache with TTL"""
        try:
            conn = sqlite3.connect(self.db_path)
            cursor = conn.cursor()
            cursor.execute('''
                INSERT OR REPLACE INTO cache (key, value, timestamp, ttl)
                VALUES (?, ?, ?, ?)
            ''', (key, json.dumps(value, ensure_ascii=False), time.time(), ttl))
            conn.commit()
            conn.close()
        except Exception as e:
            self.logger.error(f"Failed to write cache ({key}): {e}")

    def clear_expired(self):
        """Clear expired cache entries"""
        conn = sqlite3.connect(self.db_path)
        cursor = conn.cursor()
        current_time = time.time()
        cursor.execute('''
            DELETE FROM cache WHERE (timestamp + ttl) < ?
        ''', (current_time,))
        deleted = cursor.rowcount
        conn.commit()
        conn.close()
        self.logger.info(f"Cleared {deleted} expired cache entries")
        return deleted


# Rate Limiting

class RateLimiter:
    """API Rate Limiter"""

    def __init__(self):
        self.last_calls: Dict[str, float] = {}
        self.delays = {
            'pubmed': 0.6,
            'europe_pmc': 0.35,
            'kegg': 0.5,
            'uniprot': 0.8,
            'gtex': 1.2,
            'hpa': 1.2,
            'gemini': 0.3,
            'qwen': 0.3,
            'deepseek': 0.3,
            'default': 2.5
        }
        self.logger = logging.getLogger(__name__)

    async def wait(self, api_name: str):
        """Wait with configured delay for API"""
        delay = self.delays.get(api_name, 0.5)
        last_call = self.last_calls.get(api_name, 0)
        elapsed = time.time() - last_call

        if elapsed < delay:
            wait_time = delay - elapsed
            await asyncio.sleep(wait_time)

        self.last_calls[api_name] = time.time()

    async def throttle(self, api_name: str, min_interval_ms: int):
        """Precise throttle control at millisecond level"""
        min_interval = min_interval_ms / 1000.0
        last_call = self.last_calls.get(api_name, 0)
        elapsed = time.time() - last_call

        if elapsed < min_interval:
            wait_time = min_interval - elapsed
            await asyncio.sleep(wait_time)

        self.last_calls[api_name] = time.time()


# Concurrency Control

class ConcurrencyManager:
    """Concurrent Task Manager"""

    def __init__(self):
        self.logger = logging.getLogger(__name__)

    async def process_batch(
        self,
        items: List[Any],
        processor: Callable,
        max_concurrent: int = 3,
        timeout: int = 30
    ) -> List[Dict]:
        """Batch process tasks with concurrency control"""
        semaphore = asyncio.Semaphore(max_concurrent)

        async def process_with_semaphore(item):
            async with semaphore:
                start_time = time.time()
                try:
                    result = await asyncio.wait_for(
                        processor(item),
                        timeout=timeout
                    )

                    elapsed = time.time() - start_time
                    if elapsed > timeout * 0.8:
                        self.logger.warning(f"Item {item} took {elapsed:.1f}s (near timeout)")

                    return {'success': True, 'result': result, 'item': item}

                except asyncio.TimeoutError:
                    self.logger.warning(f"Processing timeout: {item} (>{timeout}s)")
                    return {
                        'success': False,
                        'error': f'Timeout after {timeout}s',
                        'item': item
                    }

                except Exception as e:
                    self.logger.error(f"Processing failed ({item}): {e}")
                    return {
                        'success': False,
                        'error': str(e),
                        'item': item
                    }

        tasks = [process_with_semaphore(item) for item in items]
        return await asyncio.gather(*tasks)


# Data Validation

class DataValidator:
    """Data Validator"""

    @staticmethod
    def validate_email(email: str) -> bool:
        """Validate email format"""
        import re
        pattern = r'^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$'
        return bool(re.match(pattern, email))

    @staticmethod
    def validate_inputs(inputs: Dict) -> Tuple[bool, str]:
        """Validate input parameters"""
        if not inputs.get('pollutant'):
            return False, 'Please enter pollutant name'

        if not inputs.get('proteins') or len(inputs['proteins']) == 0:
            return False, 'Please enter at least one candidate protein'

        if not inputs.get('target_tissue'):
            return False, 'Please select target tissue'

        if not inputs.get('email') or not DataValidator.validate_email(inputs['email']):
            return False, 'Please enter a valid email address'

        if not inputs.get('api_key'):
            return False, 'Please enter API key'

        if inputs.get('literature_limit', 0) <= 0:
            return False, 'Literature limit must be greater than 0'

        if len(inputs.get('selected_experts', [])) < 3:
            return False, 'Please select at least 3 experts'

        return True, ''


# Session Management

class SessionManager:
    """Session Save/Load Manager"""

    def __init__(self, session_dir: str = 'sessions'):
        self.session_dir = Path(session_dir)
        self.session_dir.mkdir(exist_ok=True)
        self.logger = logging.getLogger(__name__)

    def save_session(self, session_data: Dict, session_name: str = None) -> str:
        """Save analysis session"""
        if session_name is None:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            pollutant = session_data.get('inputs', {}).get('pollutant', 'unknown')
            session_name = f"{pollutant}_{timestamp}"

        session_name = "".join(c for c in session_name if c.isalnum() or c in (' ', '_', '-')).strip()
        filepath = self.session_dir / f"{session_name}.pkl"

        try:
            with open(filepath, 'wb') as f:
                pickle.dump(session_data, f)

            self.logger.info(f"Session saved: {filepath}")
            return str(filepath)
        except Exception as e:
            self.logger.error(f"Failed to save session: {e}")
            raise

    def load_session(self, filepath: str) -> Dict:
        """Load analysis session"""
        try:
            with open(filepath, 'rb') as f:
                session_data = pickle.load(f)

            self.logger.info(f"Session loaded: {filepath}")
            return session_data
        except Exception as e:
            self.logger.error(f"Failed to load session: {e}")
            raise

    def list_sessions(self) -> List[Dict]:
        """List all saved sessions"""
        sessions = []
        for filepath in self.session_dir.glob('*.pkl'):
            try:
                stat = filepath.stat()
                sessions.append({
                    'name': filepath.stem,
                    'path': str(filepath),
                    'size': stat.st_size,
                    'modified': datetime.fromtimestamp(stat.st_mtime)
                })
            except Exception as e:
                self.logger.warning(f"Failed to read session info ({filepath}): {e}")

        sessions.sort(key=lambda x: x['modified'], reverse=True)
        return sessions

    def delete_session(self, filepath: str):
        """Delete session file"""
        try:
            Path(filepath).unlink()
            self.logger.info(f"Session deleted: {filepath}")
        except Exception as e:
            self.logger.error(f"Failed to delete session: {e}")
            raise


# Research Notebook

class ResearchNotebook:
    """Research Notebook - Record key findings"""

    def __init__(self, session_name: str):
        self.notebook_data = {
            'session_name': session_name,
            'created_at': datetime.now().isoformat(),
            'sections': {
                'pollutant_background': {
                    'summary': '',
                    'key_findings': [],
                    'important_papers': []
                },
                'protein_analysis': {},
                'mechanistic_insights': [],
                'critical_questions': [],
                'synthesis_notes': '',
                'database_queries': {}
            }
        }
        self.logger = logging.getLogger(__name__)

    def add_pollutant_background_note(self, extracted_info: Dict):
        """Add pollutant background note"""
        pmid = extracted_info.get('pmid')

        if not pmid:
            return

        paper_note = {
            'pmid': pmid,
            'title': extracted_info.get('title', ''),
            'year': extracted_info.get('year', ''),
            'key_points': []
        }

        if extracted_info.get('toxic_effects'):
            paper_note['key_points'].append(f"Toxic Effects: {', '.join(extracted_info['toxic_effects'])}")

        if extracted_info.get('molecular_mechanisms'):
            paper_note['key_points'].append(f"Mechanisms: {', '.join(extracted_info['molecular_mechanisms'])}")

        if extracted_info.get('key_pathways'):
            paper_note['key_points'].append(f"Pathways: {', '.join(extracted_info['key_pathways'])}")

        if extracted_info.get('key_findings'):
            for finding in extracted_info['key_findings']:
                paper_note['key_points'].append(f"• {finding}")

        self.notebook_data['sections']['pollutant_background']['important_papers'].append(paper_note)

        self.logger.info(f"Added pollutant background note: PMID {pmid}")

    # --- MODIFICATION 1: Replace add_protein_analysis_note ---
    def add_protein_analysis_note(self, protein: str, extracted_info: Dict):
        """Add protein analysis note"""
        if protein not in self.notebook_data['sections']['protein_analysis']:
            self.notebook_data['sections']['protein_analysis'][protein] = {
                'function_summary': '',
                'interaction_evidence': [],
                'key_papers': []
            }

        pmid = extracted_info.get('pmid')
        if not pmid:
            self.logger.warning(f"Skipping protein note for {protein} due to missing PMID.")
            return

        extraction_type = extracted_info.get('extraction_type', 'unknown')

        paper_note = {
            'pmid': pmid,
            'title': extracted_info.get('title', ''),
            'year': extracted_info.get('year', ''),
            'type': extraction_type,
            'key_points': []
        }

        def add_point(label, data, is_list=True):
            if not data:
                return
            if is_list and isinstance(data, list):
                if data:
                    paper_note['key_points'].append(f"{label}: {'; '.join(data)}")
            elif not is_list and isinstance(data, str) and data not in ["Not specified", "N/A", "Unclear", "", "Not mentioned"]:
                paper_note['key_points'].append(f"{label}: {data}")
        
        add_point("Primary Functions", extracted_info.get('primary_functions'))
        add_point("Pathways", extracted_info.get('pathways'))
        add_point("Localization", extracted_info.get('localization'))
        add_point("Molecular Interactions", extracted_info.get('molecular_interactions'))
        add_point("Functional Link to Phenotype", extracted_info.get('functional_link_to_phenotype'))
        add_point("Pollutant Impact on Function", extracted_info.get('pollutant_impact_on_function'))
        add_point("Key Phenotype Results Data", extracted_info.get('key_phenotype_results_data'))
        add_point("Author Conclusions on Phenotype", extracted_info.get('author_conclusions_on_phenotype'))
        add_point("Evidence Strength", extracted_info.get('evidence_strength'), is_list=False)
        add_point("Interaction Type", extracted_info.get('interaction_type'), is_list=False)
        add_point("Functional Impact", extracted_info.get('functional_impact'), is_list=False)

        if not paper_note['key_points'] and extracted_info.get('key_findings'):
             add_point("Key Findings", extracted_info.get('key_findings'))

        if not paper_note['key_points']:
            paper_note['key_points'].append("No specific key points were extracted by the AI for this paper.")


        self.notebook_data['sections']['protein_analysis'][protein]['key_papers'].append(paper_note)
        self.logger.info(f"Added unified protein analysis note: {protein} - PMID {pmid}")
    # --- END MODIFICATION 1 ---

    def add_database_query_record(self, protein: str, database_results: Dict):
        """Record database query results"""
        if protein not in self.notebook_data['sections']['database_queries']:
            self.notebook_data['sections']['database_queries'][protein] = {
                'query_timestamp': datetime.now().isoformat(),
                'kegg': {},
                'uniprot': {},
                'expression': {},
                'data_availability': {}
            }

        record = self.notebook_data['sections']['database_queries'][protein]

        
        kegg_func_data = database_results.get('functional_data', {}).get('kegg_pathways', {})
        
        if kegg_func_data: # Check if the processed functional data exists
            pathway_count = kegg_func_data.get('pathway_count', 0) # Get correct count
            record['kegg'] = {
                'gene_id': kegg_func_data.get('gene_id'),
                'pathway_count': pathway_count, # Use correct count
                'pathways': kegg_func_data.get('pathways', []), # <--- FIX: Was []
                'categorized': kegg_func_data.get('categorized_pathways', {})
            }
            self.logger.info(f"Recorded KEGG data for {protein}: {pathway_count} pathways") # Log correct count
        elif database_results.get('raw_data', {}).get('kegg'):
             self.logger.warning(f"Recorded KEGG data for {protein}: 0 pathways (functional data missing, raw data present)")
        
        # Process UniProt data from functional_data
        uniprot_func_data = database_results.get('functional_data', {}).get('uniprot_functions', {})
        
        if uniprot_func_data: # Check if processed functional data exists
            functions = uniprot_func_data.get('functions', [])
            record['uniprot'] = {
                'accession': uniprot_func_data.get('accession'),
                'protein_name': uniprot_func_data.get('protein_name'),
                'functions': functions, # <--- FIX: Was []
                'binding_sites': uniprot_func_data.get('domains', []) # <--- FIX: Was [] (and map 'domains' to 'binding_sites')
            }
            self.logger.info(f"Recorded UniProt data for {protein}: {len(functions)} functions")
        elif database_results.get('raw_data', {}).get('uniprot'):
             self.logger.warning(f"Recorded UniProt data for {protein}: 0 functions (functional data missing, raw data present)")

        record['data_availability'] = database_results.get('data_availability', {})
        self.logger.info(f"Database query recorded for {protein}")

    def add_expression_data(self, protein: str, expression_results: Dict):
        """Record expression data"""
        if protein not in self.notebook_data['sections']['database_queries']:
            self.notebook_data['sections']['database_queries'][protein] = {
                'query_timestamp': datetime.now().isoformat(),
                'kegg': {},
                'uniprot': {},
                'expression': {},
                'data_availability': {}
            }

        record = self.notebook_data['sections']['database_queries'][protein]

        if expression_results:
            record['expression'] = {
                'tissue': expression_results.get('tissue'),
                'tpm': expression_results.get('tpm'),
                'expression_level': expression_results.get('expressionLevel'),
                'percentile': expression_results.get('percentile'),
                'source': expression_results.get('source'),
                'query_timestamp': datetime.now().isoformat()
            }
            self.logger.info(f"Recorded expression data for {protein}: TPM={expression_results.get('tpm')}")

    def add_database_query_failure(self, protein: str, error_message: str):
        """Record a failed database query"""
        if protein not in self.notebook_data['sections']['database_queries']:
            self.notebook_data['sections']['database_queries'][protein] = {}

        self.notebook_data['sections']['database_queries'][protein]['error'] = {
            'message': error_message,
            'timestamp': datetime.now().isoformat()
        }
        self.logger.warning(f"Recorded query failure for {protein}: {error_message}")

    def add_mechanistic_insight(self, insight: str, supporting_pmids: List[str], confidence: str):
        """Add mechanistic insight"""
        self.notebook_data['sections']['mechanistic_insights'].append({
            'insight': insight,
            'supporting_pmids': supporting_pmids,
            'confidence': confidence,
            'timestamp': datetime.now().isoformat()
        })

    def add_critical_questions(self, questions: List[str], source: str):
        """Add critical questions"""
        for q in questions:
            self.notebook_data['sections']['critical_questions'].append({
                'question': q,
                'source': source,
                'timestamp': datetime.now().isoformat()
            })

    def update_synthesis_notes(self, notes: str):
        """Update synthesis notes"""
        # self.notebook_data['sections']['synthesis_notes'] = notes
        if not self.notebook_data['sections']['synthesis_notes']:
            self.notebook_data['sections']['synthesis_notes'] = notes
        else:
            self.notebook_data['sections']['synthesis_notes'] += f"\n\n---\n\n{notes}"

    def export_to_markdown(self, output_path: str):
        """Export to Markdown"""
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(f"# Research Notebook: {self.notebook_data['session_name']}\n\n")
            f.write(f"**Created**: {self.notebook_data['created_at']}\n\n")
            f.write("---\n\n")

            f.write("## Pollutant Background\n\n")
            bg = self.notebook_data['sections']['pollutant_background']

            if bg.get('summary'):
                f.write(f"{bg['summary']}\n\n")

            if bg.get('important_papers'):
                f.write("### Key Papers\n\n")
                for paper in bg['important_papers']:
                    f.write(f"#### PMID: {paper['pmid']} - {paper.get('title', 'N/A')}\n")
                    f.write(f"*Year: {paper.get('year', 'N/A')}*\n\n")
                    for point in paper.get('key_points', []):
                        f.write(f"- {point}\n")
                    f.write("\n")

            # --- MODIFICATION 2: Replace Protein Analysis export logic ---
            f.write("## Protein Analysis\n\n")
            proteins = self.notebook_data['sections']['protein_analysis']

            for protein, data in proteins.items():
                f.write(f"### {protein}\n\n")

                if data.get('function_summary'):
                    f.write(f"**Function Summary**: {data['function_summary']}\n\n")

                if data.get('key_papers'):
                    f.write("### Key Papers\n\n")
                    for paper in data['key_papers']:
                        f.write(f"#### PMID: {paper['pmid']} - {paper.get('title', 'N/A')}\n")
                        f.write(f"*Year: {paper.get('year', 'N/A')} | Type: {paper.get('type', 'N/A')}*\n\n")
                        
                        for point in paper.get('key_points', []):
                            f.write(f"- {point}\n")
                        f.write("\n")
            # --- END MODIFICATION 2 ---

            db_queries = self.notebook_data['sections']['database_queries']
            if db_queries:
                f.write("## Database Query Records\n\n")
                for protein, data in db_queries.items():
                    f.write(f"### {protein}\n\n")
                    f.write(f"*Query Time: {data.get('query_timestamp', 'N/A')}*\n\n")

                    if data.get('data_availability'):
                        f.write("**Database Query Summary**:\n\n")
                        avail = data.get('data_availability', {})
                        
                        # Predefined list of all databases that should have been queried
                        db_checklist = ['kegg', 'uniprot', 'hpa', 'gtex']
                        
                        for db_name in db_checklist:
                            status = "Found" if avail.get(db_name, False) else "Not Found / Not Queried"
                            details = ""
                            if status == "Found":
                                if db_name == 'kegg' and data.get('kegg'):
                                    details = f"({data['kegg'].get('pathway_count', 0)} pathways)"
                                elif db_name == 'uniprot' and data.get('uniprot'):
                                    details = f"({len(data['uniprot'].get('functions', []))} functions)"
                                elif (db_name == 'hpa' or db_name == 'gtex') and data.get('expression'):
                                    # Use 'expression' field as HPA/GTEx data is consolidated there
                                    details = f"(TPM: {data['expression'].get('tpm', 'N/A')})"
                            
                            f.write(f"- **{db_name.upper()}**: {status} {details}\n")
                        f.write("\n")

                    if data.get('kegg') and data['kegg'].get('pathway_count', 0) > 0:
                        f.write("**KEGG Details**:\n")
                        f.write(f"- Gene ID: {data['kegg'].get('gene_id', 'N/A')}\n")
                        f.write(f"- Total Pathways: {data['kegg'].get('pathway_count', 0)}\n")
                        categorized = data['kegg'].get('categorized', {})
                        if categorized:
                            f.write("- Top Pathways by Category:\n")
                            for category, pathways in categorized.items():
                                if pathways:
                                    f.write(f"  - **{category.title()}** ({len(pathways)} total):\n")
                                    for p in pathways[:3]: # Show top 3 per category
                                        f.write(f"    - {p.get('pathway_name', 'N/A')}\n")
                        f.write("\n")

                    if data.get('uniprot'):
                        uniprot = data['uniprot']
                        f.write("**UniProt Details**:\n")
                        f.write(f"- Accession: {uniprot.get('accession', 'N/A')}\n")
                        f.write(f"- Protein Name: {uniprot.get('protein_name', 'N/A')}\n")
                        if uniprot.get('functions'):
                            # [FIX 1 & 2] Changed title, removed [:3] slice, removed [:100]...
                            f.write(f"- Functions ({len(uniprot['functions'])} total):\n")
                            for func in uniprot['functions']:
                                f.write(f"  - {func}\n")
                        if uniprot.get('binding_sites'):
                            # [FIX 2] Changed title, removed [:3] slice
                            f.write(f"- Binding Sites/Domains ({len(uniprot['binding_sites'])} total):\n")
                            for site in uniprot['binding_sites']:
                                f.write(f"  - {site}\n")
                        f.write("\n")

                    if data.get('expression'):
                        expr = data['expression']
                        f.write("**Expression Data Details**:\n")
                        f.write(f"- Tissue: {expr.get('tissue', 'N/A')}\n")
                        f.write(f"- TPM: {expr.get('tpm', 'N/A')}\n")
                        f.write(f"- Level: {expr.get('expression_level', 'N/A')}\n")
                        f.write(f"- Percentile: {expr.get('percentile', 'N/A')}%\n")
                        f.write(f"- Source: {expr.get('source', 'N/A')}\n")
                        f.write("\n")
                    
                    if data.get('error'):
                        f.write(f"**Query Error**: {data['error'].get('message', 'Unknown error')}\n\n")

            insights = self.notebook_data['sections']['mechanistic_insights']
            if insights:
                f.write("## Mechanistic Insights\n\n")
                for idx, insight in enumerate(insights, 1):
                    f.write(f"{idx}. **{insight['insight']}**\n")
                    f.write(f"   - Supporting PMIDs: {', '.join(insight['supporting_pmids'])}\n")
                    f.write(f"   - Confidence: {insight['confidence']}\n\n")

            questions = self.notebook_data['sections']['critical_questions']
            if questions:
                f.write("## Critical Questions & Gaps\n\n")
                for q in questions:
                    f.write(f"- {q['question']} *(from {q['source']})*\n")
                f.write("\n")

            if self.notebook_data['sections'].get('synthesis_notes'):
                f.write("## Synthesis Notes (PI Summaries)\n\n")
                f.write(f"{self.notebook_data['sections']['synthesis_notes']}\n")

        self.logger.info(f"Notebook exported to: {output_path}")

    def export_to_json(self, output_path: str):
        """Export to JSON"""
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(self.notebook_data, f, ensure_ascii=False, indent=2)

        self.logger.info(f"Notebook JSON exported to: {output_path}")

    def export_to_txt(self, output_path: str):
        """Export to plain text format"""
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write("="*80 + "\n")
            f.write(f"RESEARCH NOTEBOOK: {self.notebook_data['session_name']}\n")
            f.write("="*80 + "\n\n")
            f.write(f"Created: {self.notebook_data['created_at']}\n\n")
            f.write("="*80 + "\n\n")

            f.write("POLLUTANT BACKGROUND\n")
            f.write("-"*80 + "\n\n")
            bg = self.notebook_data['sections']['pollutant_background']
            if bg.get('summary'):
                f.write(f"Summary:\n{bg['summary']}\n\n")
            if bg.get('important_papers'):
                f.write("Key Papers:\n\n")
                for paper in bg['important_papers']:
                    f.write(f"  PMID: {paper['pmid']} - {paper.get('title', 'N/A')}\n")
                    f.write(f"  Year: {paper.get('year', 'N/A')}\n")
                    for point in paper.get('key_points', []):
                        f.write(f"    • {point}\n")
                    f.write("\n")
            f.write("\n" + "="*80 + "\n\n")

            f.write("PROTEIN ANALYSIS\n")
            f.write("-"*80 + "\n\n")
            proteins = self.notebook_data['sections']['protein_analysis']
            for protein, data in proteins.items():
                f.write(f"[{protein}]\n")
                f.write("-"*40 + "\n\n")
                if data.get('function_summary'):
                    f.write(f"Function: {data['function_summary']}\n\n")
                
                # Modified TXT export to match new structure
                if data.get('key_papers'):
                    f.write("Key Papers:\n\n")
                    for paper in data['key_papers']:
                        f.write(f"  PMID: {paper['pmid']} - {paper.get('title', 'N/A')}\n")
                        f.write(f"  Year: {paper.get('year', 'N/A')} | Type: {paper.get('type', 'N/A')}\n")
                        for point in paper.get('key_points', []):
                            f.write(f"    • {point}\n")
                        f.write("\n")
                
            f.write("\n" + "="*80 + "\n\n")
            
            f.write("DATABASE QUERY RECORDS\n")
            f.write("-"*80 + "\n\n")
            db_queries = self.notebook_data['sections']['database_queries']
            for protein, data in db_queries.items():
                f.write(f"[{protein}]\n")
                f.write(f"  Query Time: {data.get('query_timestamp', 'N/A')}\n\n")

                if data.get('data_availability'):
                    f.write("  Database Query Summary:\n")
                    avail = data.get('data_availability', {})
                    db_checklist = ['kegg', 'uniprot', 'hpa', 'gtex']
                    for db_name in db_checklist:
                        status = "Found" if avail.get(db_name, False) else "Not Found / Not Queried"
                        f.write(f"    - {db_name.upper()}: {status}\n")
                    f.write("\n")

                if data.get('kegg') and data['kegg'].get('pathway_count', 0) > 0:
                    f.write("  KEGG Details:\n")
                    f.write(f"    - Gene ID: {data['kegg'].get('gene_id', 'N/A')}\n")
                    f.write(f"    - Total Pathways: {data['kegg'].get('pathway_count', 0)}\n")
                    categorized = data['kegg'].get('categorized', {})
                    if categorized:
                        f.write("    - Top Pathways by Category:\n")
                        for category, pathways in categorized.items():
                            if pathways:
                                f.write(f"      - {category.title()} ({len(pathways)} total):\n")
                                for p in pathways[:2]: # Show top 2
                                    f.write(f"        - {p.get('pathway_name', 'N/A')}\n")
                    f.write("\n")
                
                if data.get('uniprot'):
                    uniprot = data['uniprot']
                    f.write("  UniProt Details:\n")
                    f.write(f"    - Accession: {uniprot.get('accession', 'N/A')}\n")
                    f.write(f"    - Protein Name: {uniprot.get('protein_name', 'N/A')}\n")
                    if uniprot.get('functions'):
                        # [FIX 1 & 2] Changed title, removed [:2] slice, removed [:100]...
                        f.write(f"    - Functions ({len(uniprot['functions'])} total):\n")
                        for func in uniprot['functions']:
                            f.write(f"      - {func}\n")
                    if uniprot.get('binding_sites'):
                        # [FIX 2] Changed title, removed [:2] slice
                        f.write(f"    - Binding Sites ({len(uniprot['binding_sites'])} total):\n")
                        for site in uniprot['binding_sites']:
                            f.write(f"      - {site}\n")
                    f.write("\n")

                if data.get('expression'):
                    expr = data['expression']
                    f.write("  Expression Data Details:\n")
                    f.write(f"    - Tissue: {expr.get('tissue', 'N/A')}, TPM: {expr.get('tpm', 'N/A')}, Level: {expr.get('expression_level', 'N/A')}\n\n")
                
                if data.get('error'):
                    f.write(f"  Query Error: {data['error'].get('message', 'Unknown error')}\n\n")

            f.write("\n" + "="*80 + "\n\n")

            insights = self.notebook_data['sections']['mechanistic_insights']
            if insights:
                f.write("MECHANISTIC INSIGHTS\n")
                f.write("-"*80 + "\n\n")
                for idx, insight in enumerate(insights, 1):
                    f.write(f"{idx}. {insight['insight']}\n")
                    f.write(f"   Supporting PMIDs: {', '.join(insight['supporting_pmids'])}\n")
                    f.write(f"   Confidence: {insight['confidence']}\n\n")
                f.write("="*80 + "\n\n")

            questions = self.notebook_data['sections']['critical_questions']
            if questions:
                f.write("CRITICAL QUESTIONS & GAPS\n")
                f.write("-"*80 + "\n\n")
                for q in questions:
                    f.write(f"• {q['question']}\n")
                    f.write(f"  (from {q['source']})\n\n")
                f.write("="*80 + "\n\n")

            if self.notebook_data['sections'].get('synthesis_notes'):
                f.write("SYNTHESIS NOTES (PI SUMMARIES)\n")
                f.write("-"*80 + "\n\n")
                f.write(f"{self.notebook_data['sections']['synthesis_notes']}\n\n")
                f.write("="*80 + "\n")

        self.logger.info(f"Notebook exported to TXT: {output_path}")


# Literature Manager

logger = logging.getLogger(__name__)


class LiteratureManager:
    """Literature Manager"""

    def __init__(self, base_dir: str = 'literature_collection'):
        self.base_dir = Path(base_dir)
        self.base_dir.mkdir(exist_ok=True)

    def create_collection(
        self,
        session_data: Dict,
        collection_name: Optional[str] = None
    ) -> str:
        """Create literature collection"""
        if collection_name is None:
            pollutant = session_data['inputs']['pollutant']
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            collection_name = f"{pollutant}_{timestamp}"

        collection_name = sanitize_filename(collection_name)

        collection_path = self.base_dir / collection_name
        collection_path.mkdir(exist_ok=True)

        (collection_path / 'abstracts').mkdir(exist_ok=True)
        (collection_path / 'fulltext').mkdir(exist_ok=True)
        (collection_path / 'metadata').mkdir(exist_ok=True)
        (collection_path / 'pdfs').mkdir(exist_ok=True)

        logger.info(f"Created literature collection: {collection_path}")

        return str(collection_path)

    def export_abstracts_to_word(
        self,
        literature_data: Dict,
        output_path: str,
        session_inputs: Dict
    ):
        """Export abstracts to Word document"""
        logger.info("Starting abstract export to Word")

        doc = Document()

        self._setup_document_styles(doc)
        self._add_title_page(doc, session_inputs)
        self._add_table_of_contents(doc, literature_data)

        for protein, data in literature_data.items():
            articles = data.get('articles', [])

            if not articles:
                continue

            heading = doc.add_heading(f'{protein}', level=1)
            heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

            stats = doc.add_paragraph()
            stats.add_run(f'Total {len(articles)} articles\n').bold = True

            quality_stats = data.get('quality_metrics', {})
            if quality_stats:
                stats.add_run(
                    f"Avg relevance: {quality_stats.get('avg_relevance_score', 0):.1f}/10  |  "
                    f"High quality: {quality_stats.get('high_quality_count', 0)}  |  "
                    f"Full-text: {quality_stats.get('has_fulltext_count', 0)}\n"
                )

            stats.add_run('-' * 80)

            for idx, article in enumerate(articles, 1):
                self._add_article_to_doc(doc, article, idx)

            doc.add_page_break()

        self._add_appendix(doc, literature_data)

        doc.save(output_path)
        logger.info(f"Abstracts exported to: {output_path}")

    def _setup_document_styles(self, doc: Document):
        """Setup document styles"""
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Times New Roman'
        font.size = Pt(11)

        from docx.oxml.ns import qn
        style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

    def _add_title_page(self, doc: Document, inputs: Dict):
        """Add title page"""
        title = doc.add_heading('Literature Search Summary', level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()

        info = doc.add_paragraph()
        info.alignment = WD_ALIGN_PARAGRAPH.CENTER

        info.add_run('Research Topic\n').bold = True
        info.add_run(f"Pollutant: {inputs.get('pollutant', 'N/A')}\n")
        info.add_run(f"Target Tissue: {inputs.get('target_tissue', 'N/A')}\n")
        info.add_run(f"Candidate Proteins: {', '.join(inputs.get('proteins', []))}\n")

        doc.add_paragraph()

        timestamp = doc.add_paragraph()
        timestamp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        timestamp.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
        timestamp.add_run(f"System: ToxTarget-AI v11.0")

        doc.add_page_break()

    def _add_table_of_contents(self, doc: Document, literature_data: Dict):
        """Add table of contents"""
        doc.add_heading('Table of Contents', level=1)

        total_count = 0
        for protein, data in literature_data.items():
            count = len(data.get('articles', []))
            total_count += count

            p = doc.add_paragraph(style='List Number')
            p.add_run(f"{protein} ").bold = True
            p.add_run(f"({count} articles)")

        doc.add_paragraph()
        summary = doc.add_paragraph()
        summary.add_run(f"Total: {total_count} articles").bold = True

        doc.add_page_break()

    def _add_article_to_doc(self, doc: Document, article: Dict, index: int):
        """Add single article to document"""
        if not article or not isinstance(article, dict):
            logger.warning(f"Skipping invalid article at index {index}")
            return

        title = article.get('title', 'No title')
        if not title or title == 'No title':
            logger.warning(f"Article {index} has no title, using placeholder")
            title = f"Untitled Article {index}"

        heading = doc.add_heading(f'[{index}] {title}', level=2)

        meta = doc.add_paragraph()

        pmid = article.get('pmid', 'N/A')
        meta.add_run(f"PMID: {pmid}  |  ").font.color.rgb = RGBColor(0, 102, 204)

        year = article.get('year', 'N/A')
        meta.add_run(f"Year: {year}  |  ")

        citations = article.get('citationCount', 0)
        meta.add_run(f"Citations: {citations}  |  ")

        score = article.get('ai_relevance_score', 0)
        if score >= 7:
            score_text = meta.add_run(f"High Relevance ({score:.1f}/10)  |  ")
            score_text.font.color.rgb = RGBColor(255, 0, 0)
            score_text.bold = True
        else:
            meta.add_run(f"Relevance: {score:.1f}/10  |  ")

        strategy = article.get('search_strategy', 'N/A')
        meta.add_run(f"Strategy: {strategy} ({article.get('source', 'N/A')})")

        if article.get('pmcid'):
            pmc_marker = meta.add_run(f"  |  PMC{article['pmcid']}")
            pmc_marker.font.color.rgb = RGBColor(0, 153, 0)

        if article.get('isOpenAccess'):
            oa_marker = meta.add_run("  |  Open Access")
            oa_marker.font.color.rgb = RGBColor(0, 153, 0)
            oa_marker.bold = True

        if article.get('full_text'):
            fulltext_marker = meta.add_run("  |  Full-text available")
            fulltext_marker.font.color.rgb = RGBColor(0, 153, 0)
            fulltext_marker.bold = True

        doc.add_paragraph()

        abstract_heading = doc.add_paragraph()
        abstract_heading.add_run('Abstract:').bold = True

        abstract_text = article.get('abstract', 'No abstract')
        abstract_para = doc.add_paragraph(abstract_text)
        abstract_para.paragraph_format.left_indent = Inches(0.5)

        if article.get('ai_relevance_reason'):
            doc.add_paragraph()
            reason_heading = doc.add_paragraph()
            reason_heading.add_run('AI Analysis:').bold = True

            reason_para = doc.add_paragraph(article['ai_relevance_reason'])
            reason_para.paragraph_format.left_indent = Inches(0.5)
            reason_para.runs[0].font.color.rgb = RGBColor(102, 102, 102)

        doc.add_paragraph()
        link_para = doc.add_paragraph()
        doi = article.get('doi')
        link_url = ""
        if pmid:
            link_url = f'https://pubmed.ncbi.nlm.nih.gov/{pmid}/'
        elif doi:
            link_url = f'http://doi.org/{doi}'

        if link_url:
            link_para.add_run('Link: ')
            link_para.add_run(link_url).font.color.rgb = RGBColor(0, 102, 204)

        doc.add_paragraph('─' * 100)

    def _add_appendix(self, doc: Document, literature_data: Dict):
        """Add appendix"""
        doc.add_heading('Appendix: Search Statistics', level=1)

        strategy_counts = {}
        total_fulltext = 0

        for protein, data in literature_data.items():
            for article in data.get('articles', []):
                strategy = article.get('search_strategy', 'Unknown')
                strategy_counts[strategy] = strategy_counts.get(strategy, 0) + 1

                if article.get('full_text'):
                    total_fulltext += 1

        table = doc.add_table(rows=1, cols=2)
        table.style = 'Light Grid Accent 1'

        header_cells = table.rows[0].cells
        header_cells[0].text = 'Search Strategy'
        header_cells[1].text = 'Article Count'

        for strategy, count in sorted(strategy_counts.items(), key=lambda x: x[1], reverse=True):
            row_cells = table.add_row().cells
            row_cells[0].text = strategy
            row_cells[1].text = str(count)

        doc.add_paragraph()

        summary = doc.add_paragraph()
        summary.add_run(f"Total full-text downloads: {total_fulltext} articles").bold = True

    def _generate_filename(self, pmid: str, title: str, ext: str) -> str:
        """Generate filename"""
        clean_title = sanitize_filename(title)[:50]
        clean_title = ''.join(c for c in clean_title if c.isalnum() or c in (' ', '-', '_'))
        clean_title = clean_title.replace(' ', '_')

        return f"PMID{pmid}_{clean_title}.{ext}"


# Expression Analyzer

class ExpressionAnalyzer:
    """Expression Analyzer based on GTEx v8"""

    def __init__(self):
        self.tissue_mapping = {
            'Liver': 'Liver',
            'Kidney': 'Kidney',
            'Brain': 'Brain',
            'Heart': 'Heart',
            'Lung': 'Lung',
            'Blood': 'Blood'
        }

        self.reference_distributions = {
            'Liver': {
                'percentiles': {
                    0: 0.0, 10: 0.5, 25: 1.8, 50: 5.2, 75: 18.5,
                    90: 62.3, 95: 120.5, 99: 450.2
                },
                'expression_thresholds': {
                    'not_detected': 0.5,
                    'low': 1.0,
                    'medium': 10.0,
                    'high': 50.0,
                    'very_high': 100.0
                }
            },
            'Kidney': {
                'percentiles': {
                    0: 0.0, 10: 0.4, 25: 1.5, 50: 4.8, 75: 16.2,
                    90: 55.8, 95: 108.3, 99: 380.5
                },
                'expression_thresholds': {
                    'not_detected': 0.5,
                    'low': 1.0,
                    'medium': 10.0,
                    'high': 45.0,
                    'very_high': 95.0
                }
            }
        }

        logger.info(f"Expression analyzer initialized")

    def get_tissue_name(self, tissue: str) -> str:
        """Get standard tissue name"""
        clean_tissue = tissue.split('(')[0].strip()
        standard_name = self.tissue_mapping.get(clean_tissue, 'Liver')
        return standard_name

    def calculate_percentile(self, tpm: float, tissue: str) -> float:
        """Calculate percentile"""
        tissue_name = self.get_tissue_name(tissue)

        if tissue_name not in self.reference_distributions:
            logger.warning(f"Reference not found for {tissue_name}, using Liver")
            tissue_name = 'Liver'

        percentiles = self.reference_distributions[tissue_name]['percentiles']

        if tpm <= percentiles[0]:
            return 0.0
        elif tpm >= percentiles[99]:
            return 99.9

        sorted_percentiles = sorted(percentiles.keys())

        for i in range(len(sorted_percentiles) - 1):
            p_low = sorted_percentiles[i]
            p_high = sorted_percentiles[i+1]

            tpm_low = percentiles[p_low]
            tpm_high = percentiles[p_high]

            if tpm_low <= tpm <= tpm_high:
                if tpm_high == tpm_low:
                    return float(p_low)

                ratio = (tpm - tpm_low) / (tpm_high - tpm_low)
                percentile = p_low + ratio * (p_high - p_low)
                return round(percentile, 1)

        if tpm > percentiles[sorted_percentiles[-1]]:
             return 99.9

        return 50.0

    def classify_expression_level(self, tpm: float, tissue: str) -> str:
        """Classify expression level"""
        tissue_name = self.get_tissue_name(tissue)

        if tissue_name not in self.reference_distributions:
            logger.warning(f"Threshold not found for {tissue_name}, using Liver")
            tissue_name = 'Liver'

        thresholds = self.reference_distributions[tissue_name]['expression_thresholds']

        if tpm < thresholds['not_detected']:
            return 'Not Detected'
        elif tpm < thresholds['low']:
            return 'Very Low'
        elif tpm < thresholds['medium']:
            return 'Low'
        elif tpm < thresholds['high']:
            return 'Medium'
        elif tpm < thresholds['very_high']:
            return 'High'
        else:
            return 'Very High'

    def analyze_expression(self, tpm: float, tissue: str) -> Dict:
        """Analyze expression data"""
        percentile = self.calculate_percentile(tpm, tissue)
        expression_level = self.classify_expression_level(tpm, tissue)

        return {
            'tpm': round(tpm, 2),
            'percentile': round(percentile, 1),
            'expression_level': expression_level,
            'tissue': tissue
        }


# Quick Screening Exporters

class QuickScreeningHTMLExporter:
    """Quick Screening HTML Exporter"""

    def __init__(self):
        self.logger = logging.getLogger(__name__)

    def export_quick_screening_report(
        self,
        screening_result: Dict,
        session_inputs: Dict,
        output_path: str
    ):
        """Export Quick Screening report as HTML"""

        html_content = self._generate_html(screening_result, session_inputs)

        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(html_content)

        self.logger.info(f"Quick Screening HTML exported: {output_path}")

    def _generate_html(self, result: Dict, inputs: Dict) -> str:
        """Generate HTML content"""

        top_list = result.get('top10', [])

        html = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>Quick Screening Report - {inputs.get('pollutant', 'N/A')}</title>
    <style>
        body {{
            font-family: 'Segoe UI', sans-serif;
            max-width: 1200px;
            margin: 0 auto;
            padding: 20px;
            background: #f5f5f5;
        }}
        .container {{
            background: white;
            padding: 40px;
            border-radius: 10px;
            box-shadow: 0 2px 10px rgba(0,0,0,0.1);
        }}
        h1 {{
            color: #2c3e50;
            border-bottom: 3px solid #3498db;
            padding-bottom: 15px;
        }}
        .info-box {{
            background: #e8f4f8;
            padding: 15px;
            border-radius: 5px;
            margin: 20px 0;
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
            margin: 20px 0;
        }}
        th, td {{
            padding: 12px;
            text-align: left;
            border: 1px solid #ddd;
        }}
        th {{
            background: #3498db;
            color: white;
        }}
        tr:nth-child(even) {{ background: #f8f9fa; }}
        .high-score {{ background: #d4edda !important; }}
        .medium-score {{ background: #fff3cd !important; }}
    </style>
</head>
<body>
    <div class="container">
        <h1>Quick Screening Report</h1>
        
        <div class="info-box">
            <strong>Pollutant:</strong> {inputs.get('pollutant', 'N/A')}<br>
            <strong>Target Tissue:</strong> {inputs.get('target_tissue', 'N/A')}<br>
            <strong>Total Evaluated:</strong> {result.get('total_evaluated', 0)}<br>
            <strong>Generated:</strong> {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
        </div>
        
        <h2>Top {len(top_list)} Proteins</h2>
        
        <table>
            <thead>
                <tr>
                    <th>Rank</th>
                    <th>Protein</th>
                    <th>Final Score</th>
                    <th>Round 1 Score</th>
                </tr>
            </thead>
            <tbody>
"""

        for item in top_list:
            score = item['final_score']
            score_class = 'high-score' if score >= 8 else ('medium-score' if score >= 6 else '')

            html += f"""
                <tr class="{score_class}">
                    <td><strong>#{item['rank']}</strong></td>
                    <td><strong>{item['protein']}</strong></td>
                    <td>{score:.2f}</td>
                    <td>{item.get('round1_score', 'N/A')}</td>
                </tr>
"""

        html += """
            </tbody>
        </table>
        
        <div style="margin-top: 40px; text-align: center; color: #888;">
            <p><strong>ToxTarget-AI v11.0</strong></p>
            <p>This is a preliminary screening</p>
        </div>
    </div>
</body>
</html>
"""
        return html


class QuickScreeningWordExporter:
    """Quick Screening Word Exporter"""

    def __init__(self):
        self.logger = logging.getLogger(__name__)

    def export_quick_screening_report(
        self,
        screening_result: Dict,
        session_inputs: Dict,
        output_path: str
    ):
        """Export Quick Screening report as Word"""

        doc = Document()

        title = doc.add_heading('Quick Screening Report', level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()
        info = doc.add_paragraph()
        info.add_run('Research Information\n').bold = True
        info.add_run(f"Pollutant: {session_inputs.get('pollutant', 'N/A')}\n")
        info.add_run(f"Target Tissue: {session_inputs.get('target_tissue', 'N/A')}\n")
        info.add_run(f"Total Evaluated: {screening_result.get('total_evaluated', 0)}\n")
        info.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")

        doc.add_page_break()

        doc.add_heading('Top Proteins', level=1)

        top_list = screening_result.get('top10', [])

        table = doc.add_table(rows=1, cols=4)
        table.style = 'Light Grid Accent 1'

        header_cells = table.rows[0].cells
        header_cells[0].text = 'Rank'
        header_cells[1].text = 'Protein'
        header_cells[2].text = 'Final Score'
        header_cells[3].text = 'Round 1 Score'

        for item in top_list:
            row_cells = table.add_row().cells
            row_cells[0].text = f"#{item['rank']}"
            row_cells[1].text = item['protein']
            row_cells[2].text = f"{item['final_score']:.2f}"
            row_cells[3].text = str(item.get('round1_score', 'N/A'))

        doc.add_page_break()

        footer = doc.add_paragraph()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run('ToxTarget-AI v11.0\n').bold = True
        footer.add_run('This is a preliminary screening')

        doc.save(output_path)
        self.logger.info(f"Quick Screening Word exported: {output_path}")


class DetailedWordExporter:
    """Detailed Analysis Word Document Exporter"""
    
    def __init__(self):
        self.logger = logging.getLogger(__name__)
    
    def export_detailed_report(
        self,
        results: Dict,
        session_inputs: Dict,
        output_path: str
    ):
        """
        Export a comprehensive Word document with all analysis results
        
        Args:
            results: Dictionary containing all analysis results (session_data)
            session_inputs: Input parameters for the analysis
            output_path: Path where the Word document will be saved
        """
        try:
            doc = Document()
            
            # Title Page
            self._add_title_page(doc, session_inputs)
            doc.add_page_break()
            
            # Table of Contents (placeholder)
            doc.add_heading('Table of Contents', level=1)
            doc.add_paragraph('1. Research Overview')
            doc.add_paragraph('2. Quick Screening Results')
            doc.add_paragraph('3. Literature Analysis')
            doc.add_paragraph('4. Tissue Expression Analysis')
            doc.add_paragraph('5. Database Analysis')
            doc.add_paragraph('6. Function-Phenotype Analysis')
            doc.add_paragraph('7. Expert Discussion Summary')
            doc.add_paragraph('8. Final Evaluation and Conclusions')
            doc.add_page_break()
            
            # Section 1: Research Overview
            self._add_research_overview(doc, session_inputs)
            doc.add_page_break()
            
            # Section 2: Quick Screening Results
            if 'quick_screening_result' in results:
                self._add_quick_screening_section(doc, results['quick_screening_result'])
                doc.add_page_break()
            
            # Section 3: Literature Analysis
            if 'literature' in results or 'literature_stage1' in results:
                self._add_literature_section(doc, results)
                doc.add_page_break()
            
            # Section 4: Tissue Expression Analysis
            if 'tissue_expression' in results:
                self._add_tissue_expression_section(doc, results['tissue_expression'])
                doc.add_page_break()
            
            # Section 5: Database Analysis
            if 'database_analysis' in results:
                self._add_database_analysis_section(doc, results['database_analysis'])
                doc.add_page_break()
            
            # Section 6: Function-Phenotype Analysis
            if 'function_phenotype_analysis' in results:
                self._add_function_phenotype_section(doc, results['function_phenotype_analysis'])
                doc.add_page_break()
            
            # Section 7: Expert Discussion
            if 'round1_meeting' in results or 'round2_meeting' in results:
                self._add_expert_discussion_section(doc, results)
                doc.add_page_break()
            
            # Section 8: Final Evaluation
            if 'final_evaluation' in results:
                self._add_final_evaluation_section(doc, results['final_evaluation'])
            
            # Save document
            doc.save(output_path)
            self.logger.info(f"Detailed Word report exported: {output_path}")
            
        except Exception as e:
            self.logger.error(f"Failed to export detailed Word report: {e}", exc_info=True)
            raise
    
    def _add_title_page(self, doc: Document, inputs: Dict):
        """Add title page"""
        title = doc.add_heading('Toxicological Target Protein Analysis', level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph()
        
        subtitle = doc.add_heading('Comprehensive Screening Report', level=2)
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph()
        doc.add_paragraph()
        
        info = doc.add_paragraph()
        info.alignment = WD_ALIGN_PARAGRAPH.CENTER
        info.add_run(f"Pollutant: {inputs.get('pollutant', 'N/A')}\n").bold = True
        info.add_run(f"Target Tissue: {inputs.get('target_tissue', 'N/A')}\n")
        info.add_run(f"Analysis Mode: {'Quick Screening' if inputs.get('is_quick_mode', False) else 'Detailed Analysis'}\n")
        info.add_run(f"\nReport Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
    
    def _add_research_overview(self, doc: Document, inputs: Dict):
        """Add research overview section"""
        doc.add_heading('1. Research Overview', level=1)
        
        doc.add_heading('1.1 Research Parameters', level=2)
        params = doc.add_paragraph()
        params.add_run('Pollutant: ').bold = True
        params.add_run(f"{inputs.get('pollutant', 'N/A')}\n")
        params.add_run('Target Tissue: ').bold = True
        params.add_run(f"{inputs.get('target_tissue', 'N/A')}\n")
        params.add_run('Target Phenotypes: ').bold = True
        params.add_run(f"{inputs.get('target_phenotypes', 'N/A')}\n")
        params.add_run('Candidate Proteins: ').bold = True
        params.add_run(f"{inputs.get('candidate_proteins', 'N/A')}\n")
        
        if inputs.get('hypothesis'):
            doc.add_heading('1.2 Research Hypothesis', level=2)
            doc.add_paragraph(inputs['hypothesis'])
    
    def _add_quick_screening_section(self, doc: Document, quick_result: Dict):
        """Add quick screening results section"""
        doc.add_heading('2. Quick Screening Results', level=1)
        
        top_proteins = quick_result.get('top10', [])
        
        if top_proteins:
            doc.add_heading('2.1 Top Ranked Proteins', level=2)
            
            # Create table
            table = doc.add_table(rows=1, cols=4)
            table.style = 'Light Grid Accent 1'
            
            header_cells = table.rows[0].cells
            header_cells[0].text = 'Rank'
            header_cells[1].text = 'Protein'
            header_cells[2].text = 'Final Score'
            header_cells[3].text = 'Round 1 Score'
            
            for item in top_proteins:
                row_cells = table.add_row().cells
                row_cells[0].text = f"#{item.get('rank', 'N/A')}"
                row_cells[1].text = str(item.get('protein', 'N/A'))
                row_cells[2].text = f"{item.get('final_score', 0):.2f}"
                row_cells[3].text = str(item.get('round1_score', 'N/A'))
        
        doc.add_paragraph()
        summary = doc.add_paragraph()
        summary.add_run('Summary Statistics:\n').bold = True
        summary.add_run(f"Total Proteins Evaluated: {quick_result.get('total_evaluated', 0)}\n")
        summary.add_run(f"Proteins in Top 10: {len(top_proteins)}\n")
    
    def _add_literature_section(self, doc: Document, results: Dict):
        """Add literature analysis section"""
        doc.add_heading('3. Literature Analysis', level=1)
        
        if 'pollutant_background_literature' in results:
            doc.add_heading('3.1 Pollutant Background', level=2)
            bg_lit = results['pollutant_background_literature']
            if isinstance(bg_lit, dict) and 'summary' in bg_lit:
                doc.add_paragraph(bg_lit['summary'])
            elif isinstance(bg_lit, str):
                doc.add_paragraph(bg_lit)
        
        if 'literature' in results and 'selected' in results['literature']:
            doc.add_heading('3.2 Selected Literature', level=2)
            selected = results['literature']['selected']
            for protein, lit_list in selected.items():
                doc.add_heading(f'Protein: {protein}', level=3)
                if lit_list:
                    for i, lit in enumerate(lit_list, 1):
                        p = doc.add_paragraph(style='List Number')
                        if isinstance(lit, dict):
                            title = lit.get('title', 'No title')
                            pmid = lit.get('pmid', 'N/A')
                            p.add_run(f"{title} ").bold = True
                            p.add_run(f"(PMID: {pmid})")
                        else:
                            p.add_run(str(lit))
    
    def _add_tissue_expression_section(self, doc: Document, tissue_expr: Dict):
        """Add tissue expression analysis section"""
        doc.add_heading('4. Tissue Expression Analysis', level=1)
        
        if not tissue_expr:
            doc.add_paragraph('No tissue expression data available.')
            return
        
        # Create table for expression data
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Light Grid Accent 1'
        
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Protein'
        header_cells[1].text = 'TPM Value'
        header_cells[2].text = 'Percentile'
        header_cells[3].text = 'Expression Level'
        header_cells[4].text = 'Tissue'
        
        for protein, expr_data in tissue_expr.items():
            row_cells = table.add_row().cells
            row_cells[0].text = protein
            row_cells[1].text = f"{expr_data.get('tpm', 0):.2f}"
            row_cells[2].text = f"{expr_data.get('percentile', 0):.1f}"
            row_cells[3].text = expr_data.get('expression_level', 'N/A')
            row_cells[4].text = expr_data.get('tissue', 'N/A')
    
    def _add_database_analysis_section(self, doc: Document, db_analysis: Dict):
        """Add database analysis section"""
        doc.add_heading('5. Database Analysis', level=1)
        
        if not db_analysis:
            doc.add_paragraph('No database analysis available.')
            return
        
        for protein, data in db_analysis.items():
            doc.add_heading(f'5.{list(db_analysis.keys()).index(protein) + 1} {protein}', level=2)
            
            if isinstance(data, dict):
                # Add pathways
                if 'pathways' in data and data['pathways']:
                    doc.add_heading('Pathways:', level=3)
                    for pathway in data['pathways'][:10]:  # Limit to top 10
                        doc.add_paragraph(pathway, style='List Bullet')
                
                # Add GO terms
                if 'go_terms' in data and data['go_terms']:
                    doc.add_heading('GO Terms:', level=3)
                    for go_term in data['go_terms'][:10]:  # Limit to top 10
                        if isinstance(go_term, dict):
                            p = doc.add_paragraph(style='List Bullet')
                            p.add_run(f"{go_term.get('term', 'N/A')} ").bold = True
                            p.add_run(f"({go_term.get('category', 'N/A')})")
                        else:
                            doc.add_paragraph(str(go_term), style='List Bullet')
            else:
                doc.add_paragraph(str(data))
    
    def _add_function_phenotype_section(self, doc: Document, fp_analysis: Dict):
        """Add function-phenotype analysis section"""
        doc.add_heading('6. Function-Phenotype Analysis', level=1)
        
        if not fp_analysis:
            doc.add_paragraph('No function-phenotype analysis available.')
            return
        
        for protein, analysis in fp_analysis.items():
            doc.add_heading(f'6.{list(fp_analysis.keys()).index(protein) + 1} {protein}', level=2)
            
            if isinstance(analysis, dict):
                # Add analysis summary
                if 'summary' in analysis:
                    doc.add_paragraph(analysis['summary'])
                
                # Add detailed analysis
                if 'detailed_analysis' in analysis:
                    doc.add_heading('Detailed Analysis:', level=3)
                    doc.add_paragraph(analysis['detailed_analysis'])
                
                # Add relevance score if available
                if 'relevance_score' in analysis:
                    p = doc.add_paragraph()
                    p.add_run('Relevance Score: ').bold = True
                    p.add_run(f"{analysis['relevance_score']}")
            else:
                doc.add_paragraph(str(analysis))
    
    def _add_expert_discussion_section(self, doc: Document, results: Dict):
        """Add expert discussion section"""
        doc.add_heading('7. Expert Discussion Summary', level=1)
        
        # Round 1 discussion
        if 'round1_meeting' in results:
            doc.add_heading('7.1 Round 1 Discussion', level=2)
            meeting = results['round1_meeting']
            
            if isinstance(meeting, dict):
                if 'final_summary' in meeting and meeting['final_summary']:
                    doc.add_paragraph(meeting['final_summary'])
                
                if 'discussion_history' in meeting and meeting['discussion_history']:
                    doc.add_heading('Discussion Highlights:', level=3)
                    for discussion in meeting['discussion_history'][:5]:  # Limit to first 5
                        if isinstance(discussion, str):
                            doc.add_paragraph(discussion[:500] + '...' if len(discussion) > 500 else discussion)
        
        # Round 2 discussion
        if 'round2_meeting' in results:
            doc.add_heading('7.2 Round 2 Discussion', level=2)
            meeting = results['round2_meeting']
            
            if isinstance(meeting, dict) and 'final_summary' in meeting:
                doc.add_paragraph(meeting['final_summary'])
    
    def _add_final_evaluation_section(self, doc: Document, final_eval: Dict):
        """Add final evaluation and conclusions section"""
        doc.add_heading('8. Final Evaluation and Conclusions', level=1)
        
        if not final_eval:
            doc.add_paragraph('No final evaluation available.')
            return
        
        # Add mini-review if available
        if 'mini_review' in final_eval:
            doc.add_heading('8.1 Comprehensive Review', level=2)
            doc.add_paragraph(final_eval['mini_review'])
        
        # Add final rankings if available
        if 'final_rankings' in final_eval:
            doc.add_heading('8.2 Final Rankings', level=2)
            rankings = final_eval['final_rankings']
            
            if isinstance(rankings, list):
                table = doc.add_table(rows=1, cols=3)
                table.style = 'Light Grid Accent 1'
                
                header_cells = table.rows[0].cells
                header_cells[0].text = 'Rank'
                header_cells[1].text = 'Protein'
                header_cells[2].text = 'Score'
                
                for i, item in enumerate(rankings, 1):
                    row_cells = table.add_row().cells
                    row_cells[0].text = str(i)
                    if isinstance(item, dict):
                        row_cells[1].text = item.get('protein', 'N/A')
                        row_cells[2].text = str(item.get('score', 'N/A'))
                    else:
                        row_cells[1].text = str(item)
                        row_cells[2].text = 'N/A'
        
        if 'conclusions' in final_eval:
            doc.add_heading('8.3 Conclusions', level=2)
            doc.add_paragraph(final_eval['conclusions'])
        
        # Add recommendations if available
        if 'recommendations' in final_eval:
            doc.add_heading('8.4 Recommendations', level=2)
            doc.add_paragraph(final_eval['recommendations'])
        
        # Footer
        doc.add_paragraph()
        footer = doc.add_paragraph()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run('─' * 50 + '\n')
        footer.add_run('End of Report\n').bold = True
        footer.add_run('ToxTarget-AI Analysis System v11.0\n')
        footer.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")


# Database CSV Exporter

class DatabaseCSVExporter:
    """Exports detailed database query results to a CSV file"""
    
    def __init__(self):
        self.logger = logging.getLogger(__name__)

    def export_database_csv(
        self,
        database_analysis: Dict[str, Dict],
        tissue_expression: Dict[str, Dict],
        output_path: str
    ):
        """
        Merges database_analysis and tissue_expression data and exports to CSV.
        """
        if not database_analysis and not tissue_expression:
            self.logger.warning("No database or expression data to export.")
            return

        # 1. Define CSV Headers
        headers = [
            'Protein',
            'UniProt_Accession',
            'UniProt_Name',
            'UniProt_Functions',
            'UniProt_Keywords',
            'UniProt_Pathways',
            'UniProt_Catalytic_Activity',
            'UniProt_Cofactors',
            'UniProt_Domains',
            'KEGG_ID',
            'KEGG_Pathways',
            'KEGG_Diseases',
            'KEGG_Modules',
            'KEGG_Orthologs',
            'Tissue_Expression_Level',
            'Tissue_TPM',
            'Tissue_Percentile',
            'Expression_Source'
        ]

        # 2. Get all protein keys
        all_proteins = set(database_analysis.keys()) | set(tissue_expression.keys())
        
        if not all_proteins:
            self.logger.warning("Protein list is empty, cannot export CSV.")
            return

        try:
            with open(output_path, 'w', newline='', encoding='utf-8-sig') as f:
                writer = csv.DictWriter(f, fieldnames=headers)
                writer.writeheader()

                # 3. Iterate over each protein and build the row
                for protein in sorted(list(all_proteins)):
                    row = {'Protein': protein}
                    
                    db_data = database_analysis.get(protein, {})
                    expr_data = tissue_expression.get(protein, {})

                    # 3a. Flatten UniProt data
                    if db_data.get('functional_data', {}).get('uniprot_functions'):
                        u_data = db_data['functional_data']['uniprot_functions']
                        row['UniProt_Accession'] = u_data.get('accession')
                        row['UniProt_Name'] = u_data.get('protein_name')
                        row['UniProt_Functions'] = self._join_list(u_data.get('functions'))
                        row['UniProt_Keywords'] = self._join_list(u_data.get('keywords'))
                        row['UniProt_Pathways'] = self._join_list(u_data.get('pathways'))
                        row['UniProt_Catalytic_Activity'] = self._join_list(u_data.get('catalytic_activity'))
                        row['UniProt_Cofactors'] = self._join_list(u_data.get('cofactors'))
                        row['UniProt_Domains'] = self._join_list(u_data.get('domains'))

                    # 3b. Flatten KEGG data
                    if db_data.get('functional_data', {}).get('kegg_pathways'):
                        k_data = db_data['functional_data']['kegg_pathways']
                        k_raw_details = db_data.get('raw_data', {}).get('kegg', {}).get('details', {})
                        
                        row['KEGG_ID'] = k_data.get('gene_id')
                        row['KEGG_Pathways'] = self._join_list(k_data.get('pathways'), key='pathway_name')
                        # Get richer info from raw parsed data
                        row['KEGG_Diseases'] = self._join_list(k_raw_details.get('disease_associations'), key='disease_name')
                        row['KEGG_Modules'] = self._join_list(k_raw_details.get('modules'), key='module_name')
                        row['KEGG_Orthologs'] = self._join_list(k_raw_details.get('orthologs'), key='ortholog_name')

                    # 3c. Flatten Tissue Expression data
                    if expr_data:
                        row['Tissue_Expression_Level'] = expr_data.get('expression_level')
                        row['Tissue_TPM'] = expr_data.get('tpm')
                        row['Tissue_Percentile'] = expr_data.get('percentile')
                        row['Expression_Source'] = expr_data.get('source')

                    writer.writerow(row)
            
            self.logger.info(f"Successfully exported database details to: {output_path}")

        except Exception as e:
            self.logger.error(f"Failed to export database CSV: {e}", exc_info=True)
            raise

    def _join_list(self, data: Optional[List], key: Optional[str] = None) -> str:
        """Helper to join a list (or list of dicts) into a semi-colon separated string"""
        if not data:
            return ""
        
        try:
            if key:
                # List of dicts
                return "; ".join([str(item.get(key, '')) for item in data if isinstance(item, dict)])
            else:
                # Simple list
                return "; ".join([str(item) for item in data])
        except Exception as e:
            self.logger.error(f"CSV _join_list failed: {e}")
            return "[Error]"